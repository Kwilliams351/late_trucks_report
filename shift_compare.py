"""Persist truck snapshots and compare shifts without rewriting the report workbook."""
from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

import pandas as pd

from report_card_pending import CARD_STATUSES, card_status

TIMES = {"STD": "sched_dep", "ATD": "act_dep", "STA": "sched_arr", "ATA": "act_arr", "Proj ATA": "proj_arr"}
VALUES = {"TLM": "tlm", "Date": "sched_date", "Truck": "truck", "Origin": "origin", "Dest": "dest",
          "Vendor": "vendor", "AC Type": "ac_type", "Arr Rule": "arr_rule", "Dep Status": "dep_status",
          "Arr Status": "arr_status", "In Window": "in_window"}
RULES = {"threshold": "Threshold", "grace": "Grace", "early_check": "EarlyCheck", "late_check": "LateCheck",
         "min_transit": "MinTransit"}
SOURCE_FIELDS = ["Date", "Truck", "Origin", "Dest", "Vendor", "AC Type", "Arr Rule", "STD", "ATD", "STA", "ATA"]
CONTEXT_FIELDS = ["Proj ATA", "Origin Offset (h)", "Dest Offset (h)", *RULES.values()]
STATUS_FIELDS = ["Dep Status", "Arr Status", "Card status"]
COMPARE_FIELDS = SOURCE_FIELDS + CONTEXT_FIELDS + STATUS_FIELDS + ["In Window"]
SNAPSHOT_COLUMNS = ["TLM", *dict.fromkeys(COMPARE_FIELDS), "AsOf"]
METRICS = ["Trucks", "Late dep", "Late arr", "No ATD", "No ATA", "Pending", "Check date", "Ready to log"]
EVENTS = ["Became late", "No longer late (time regained)", "Pending -> ready to log", "Now overdue", "Clock only"]


def _number(value):
    return format(float(value), ".10g")


def make_snapshot(df: pd.DataFrame, settings: dict) -> pd.DataFrame:
    out = pd.DataFrame(index=df.index)
    for header, source in VALUES.items():
        if header == "Date":
            out[header] = df[source].dt.strftime("%Y-%m-%d").fillna("")
        elif header == "In Window":
            out[header] = df[source].astype(int).astype(str)
        else:
            out[header] = df[source].fillna("").astype(str)
    for header, source in TIMES.items():
        out[header] = df[source].dt.strftime("%Y-%m-%d %H:%M").fillna("")
    for header, source in (("Origin Offset (h)", "now_o"), ("Dest Offset (h)", "now_d")):
        hours = (df[source] - pd.Timestamp(settings["as_of"])).dt.total_seconds() / 3600
        out[header] = hours.map(lambda v: _number(round(v, 2)))
    out["Card status"] = card_status(df, settings["threshold"])
    for key, name in RULES.items():
        out[name] = _number(settings[key])
    out["AsOf"] = pd.Timestamp(settings["as_of"]).isoformat(sep=" ")
    return _validate(out.reindex(columns=SNAPSHOT_COLUMNS).reset_index(drop=True), settings["as_of"])


def _validate(frame, as_of):
    missing = set(SNAPSHOT_COLUMNS) - set(frame.columns)
    if missing:
        raise ValueError(f"Snapshot is missing columns: {', '.join(sorted(missing))}")
    frame = frame.fillna("").astype(str)
    if frame["TLM"].eq("").any() or frame["TLM"].duplicated().any():
        raise ValueError("Snapshots require a non-blank, unique TLM for every truck.")
    if not frame["Dep Status"].isin(["Late", "On time", "No actual", "Pending", "Check date"]).all():
        raise ValueError("Missing or invalid departure statuses; open the report in Excel, recalculate and save it.")
    if not frame["Arr Status"].isin(["Late", "On time", "No actual", "Pending", "Check date", "Not tracked"]).all():
        raise ValueError("Missing or invalid arrival statuses; open the report in Excel, recalculate and save it.")
    if not frame["Card status"].isin(["", *CARD_STATUSES]).all() or not frame["In Window"].isin(["0", "1"]).all():
        raise ValueError("Invalid cached Card status or window flag; recalculate and save the report in Excel.")
    frame.attrs["as_of"] = pd.Timestamp(as_of)
    return frame


def save_snapshot(frame, directory: Path) -> Path:
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    stamp = frame.attrs["as_of"].strftime("%Y-%m-%d_%H%M%S")
    # A rerun of the same as-of must not overwrite the previous shift's evidence.
    path = directory / f"snapshot_{stamp}_{datetime.now():%Y%m%d%H%M%S%f}.csv"
    temp = path.with_suffix(".tmp")
    frame.to_csv(temp, index=False)
    temp.replace(path)
    return path


def _csv_time(path: Path, frame) -> pd.Timestamp:
    if len(frame):
        stamps = frame["AsOf"].unique()
        if len(stamps) != 1:
            raise ValueError(f"{path}: snapshot contains multiple as-of times.")
        return pd.Timestamp(stamps[0])
    return pd.Timestamp(datetime.strptime(path.stem[9:26], "%Y-%m-%d_%H%M%S"))


def read_snapshot(path: Path) -> pd.DataFrame:
    path = Path(path)
    if path.suffix.lower() == ".csv":
        frame = pd.read_csv(path, dtype=str, keep_default_na=False)
        return _validate(frame, _csv_time(path, frame))
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"Expected a snapshot CSV or a recalculated report XLSX: {path}")
    import openpyxl
    with path.open("rb") as stream:
        wb = openpyxl.load_workbook(stream, read_only=True, data_only=True)
        try:
            if "Data" not in wb or "Settings" not in wb:
                raise ValueError(f"{path}: expected a daily report with Data and Settings tabs.")
            ws = wb["Data"]
            headers = [c.value for c in ws[3]]
            required = {"TLM", *SOURCE_FIELDS, *TIMES, "Dep Status", "Arr Status", "In Window",
                        "Origin Offset (h)", "Dest Offset (h)", "Proj Arr Delay (min)"}
            if required - set(headers):
                raise ValueError(f"{path}: Data tab is missing required snapshot columns.")
            rows = [dict(zip(headers, row)) for row in ws.iter_rows(min_row=4, values_only=True) if row[0] is not None]
            frame = pd.DataFrame(rows, columns=headers)
            settings = {}
            for name in ["AsOf", *RULES.values()]:
                definition = wb.defined_names.get(name)
                if definition is None:
                    raise ValueError(f"{path}: missing Settings name {name}.")
                sheet, address = next(definition.destinations)
                settings[name] = wb[sheet][address].value
                if settings[name] is None:
                    raise ValueError(f"{path}: missing cached value for {name}; recalculate and save in Excel.")
            for name in [*TIMES, "Date"]:
                frame[name] = pd.to_datetime(frame[name], errors="raise").dt.strftime(
                    "%Y-%m-%d" if name == "Date" else "%Y-%m-%d %H:%M").fillna("")
            if "Card status" not in frame:
                proxy = pd.DataFrame({
                    "dep_status": frame["Dep Status"], "arr_status": frame["Arr Status"],
                    "act_dep": pd.to_datetime(frame["ATD"]), "act_arr": pd.to_datetime(frame["ATA"]),
                    "arr_rule": frame["Arr Rule"],
                    "proj_arr_delay": pd.to_numeric(frame["Proj Arr Delay (min)"], errors="raise"),
                    "late": frame["Dep Status"].eq("Late") | frame["Arr Status"].eq("Late"),
                })
                frame["Card status"] = card_status(proxy, settings["Threshold"])
            for name in ["Origin Offset (h)", "Dest Offset (h)", "In Window"]:
                frame[name] = frame[name].map(_number)
            for name in RULES.values():
                frame[name] = _number(settings[name])
            frame["AsOf"] = pd.Timestamp(settings["AsOf"]).isoformat(sep=" ")
            return _validate(frame[SNAPSHOT_COLUMNS], settings["AsOf"])
        finally:
            wb.close()


def select_prior(directory: Path, as_of, vs="prior-shift"):
    if vs not in ("prior-shift", "prior-day"):
        previous = read_snapshot(Path(vs))
        if previous.attrs["as_of"] >= pd.Timestamp(as_of):
            raise ValueError("The comparison report/snapshot must be earlier than the current as-of time.")
        return previous
    as_of = pd.Timestamp(as_of)
    cutoff = as_of - pd.Timedelta(days=1) if vs == "prior-day" else as_of
    candidates = []
    for path in Path(directory).glob("snapshot_*.csv"):
        first = pd.read_csv(path, dtype=str, keep_default_na=False, nrows=1)
        stamp = _csv_time(path, first)
        if stamp < as_of and stamp <= cutoff:
            candidates.append((stamp, path.name, path))
    return read_snapshot(max(candidates)[2]) if candidates else None


def _totals(frame):
    frame = frame[frame["In Window"].eq("1")]
    return dict(zip(METRICS, [
        len(frame), int(frame["Dep Status"].eq("Late").sum()), int(frame["Arr Status"].eq("Late").sum()),
        int(frame["Dep Status"].eq("No actual").sum()),
        int((frame["Arr Status"].eq("No actual") & frame["Dep Status"].ne("No actual")).sum()),
        int(frame["Card status"].str.startswith("Pending").sum()),
        int((frame["Dep Status"].eq("Check date") | frame["Arr Status"].eq("Check date")).sum()),
        int(frame["Card status"].eq("Ready to log").sum()),
    ]))


def compare_snapshots(previous, current):
    old, new = previous.set_index("TLM"), current.set_index("TLM")
    records, counts = [], dict.fromkeys(EVENTS, 0)
    for tlm in old.index.union(new.index, sort=False):
        before = old.loc[tlm] if tlm in old.index else None
        after = new.loc[tlm] if tlm in new.index else None
        changes, events = [], []
        clock_only = False
        if before is None:
            changes = ["Added"]
        elif after is None:
            changes = ["Removed"]
        else:
            changes = [name for name in COMPARE_FIELDS if before[name] != after[name]]
            if not changes:
                continue
            late_before = before["Dep Status"] == "Late" or before["Arr Status"] == "Late"
            late_after = after["Dep Status"] == "Late" or after["Arr Status"] == "Late"
            valid_after = after["Dep Status"] in ("Late", "On time") and after["Arr Status"] in ("Late", "On time", "Not tracked")
            pending_before = before["Card status"].startswith("Pending")
            overdue_before = before["Card status"] == "Pending - arrival overdue" or before["Dep Status"] == "No actual"
            overdue_after = after["Card status"] == "Pending - arrival overdue" or after["Dep Status"] == "No actual"
            checks = [not late_before and late_after, late_before and not late_after and valid_after,
                      pending_before and after["Card status"] == "Ready to log", not overdue_before and overdue_after]
            events = [event for event, applies in zip(EVENTS, checks) if applies]
            clock_only = (
                previous.attrs["as_of"] != current.attrs["as_of"]
                and not any(name in changes for name in SOURCE_FIELDS + CONTEXT_FIELDS)
                and any(name in changes for name in STATUS_FIELDS)
            )
            if clock_only:
                events.append("Clock only")
            # Transition totals use the same cohort, not trucks entering/leaving a sliding window.
            if before["In Window"] == after["In Window"] == "1":
                for event in events:
                    counts[event] += 1
        record = {"TLM": tlm, "Changed": ", ".join(changes), "Transitions": "; ".join(events),
                  "Clock only": "Y" if clock_only else ""}
        for name in COMPARE_FIELDS:
            record["Prior " + name] = before[name] if before is not None else ""
            record["Current " + name] = after[name] if after is not None else ""
        records.append(record)
    return {"prior_as_of": previous.attrs["as_of"], "current_as_of": current.attrs["as_of"],
            "prior_totals": _totals(previous), "current_totals": _totals(current),
            "transitions": counts, "changes": records}


def write_variance(wb, comparison):
    ws = wb.add_worksheet("Shift Variance")
    title = wb.add_format({"bold": True, "font_color": "white", "bg_color": "#1F3864"})
    wrap = wb.add_format({"text_wrap": True, "valign": "top"})
    ws.write(0, 0, "Shift Variance", title)
    ws.write(1, 0, f"Prior {comparison['prior_as_of']} | Current {comparison['current_as_of']}")
    ws.write(2, 0, "Totals use each report's window; transitions use trucks in both windows. "
                   "Clock only = times, projection and rules unchanged; only the as-of clock advanced.")
    ws.write_row(4, 0, ["Metric", "Prior", "Current", "Change"], title)
    for r, metric in enumerate(METRICS, 5):
        before, after = comparison["prior_totals"][metric], comparison["current_totals"][metric]
        ws.write_row(r, 0, [metric, before, after, after - before])
    start = 6 + len(METRICS)
    ws.write_row(start, 0, ["Transition (comparable cohort)", "Trucks"], title)
    for r, event in enumerate(EVENTS, start + 1):
        ws.write_row(r, 0, [event, comparison["transitions"][event]])
    header_row = start + len(EVENTS) + 3
    headers = ["TLM", "Changed", "Transitions", "Clock only",
               *[prefix + name for name in COMPARE_FIELDS for prefix in ("Prior ", "Current ")]]
    ws.write_row(header_row, 0, headers, title)
    for row, record in enumerate(comparison["changes"], header_row + 1):
        # Snapshot text is data, never an Excel formula or an automatic URL.
        for col, header in enumerate(headers):
            ws.write_string(row, col, str(record[header]), wrap)
    last = header_row + len(comparison["changes"])
    ws.autofilter(header_row, 0, max(header_row + 1, last), len(headers) - 1)
    ws.freeze_panes(header_row + 1, 1)
    ws.set_column(0, 0, 31)
    ws.set_column(1, 2, 45)
    ws.set_column(3, len(headers) - 1, 20)
    if not comparison["changes"]:
        ws.write(header_row + 1, 0, "No truck-level changes.")


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    for command in ("snapshot", "compare"):
        cmd = sub.add_parser(command)
        cmd.add_argument("report", type=Path)
        cmd.add_argument("--snapshots", type=Path, help="default: snapshots next to the report's Output folder")
        if command == "compare":
            cmd.add_argument("--vs", default="prior-shift", help="prior-shift, prior-day, or a prior XLSX/CSV path")
            cmd.add_argument("--out", type=Path, help="standalone comparison XLSX")
    args = parser.parse_args(argv)
    current = read_snapshot(args.report)
    directory = args.snapshots or args.report.parent.parent / "snapshots"
    if args.command == "snapshot":
        print(f"Saved {save_snapshot(current, directory)}")
        return
    previous = select_prior(directory, current.attrs["as_of"], args.vs)
    if previous is None:
        parser.exit(1, "No earlier snapshot found. Build another shift or pass --vs a prior report.\n")
    path = args.out or args.report.with_name(args.report.stem + " - Shift Variance.xlsx")
    if path.resolve() == args.report.resolve():
        parser.error("--out must not overwrite the source report.")
    import xlsxwriter
    with xlsxwriter.Workbook(str(path)) as wb:
        write_variance(wb, compare_snapshots(previous, current))
    print(f"Wrote {path}")


if __name__ == "__main__":
    main()
