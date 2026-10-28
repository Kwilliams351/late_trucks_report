import os
import shutil
import sys
import tempfile
import time
import unittest
from datetime import datetime
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import late_trucks as L  # noqa: E402

COLS = ["Scheduled Departure Date", "Truck", "Vendor Code", "Vendor Name", "Origin", "Destination", "AC Type",
        "Scheduled Departure Time", "Scheduled Arrival Date", "TLM", "Actual Departure Date", "Actual Arrival Date"]

# Custom-SQL data source export layout (headers as exported from Tableau)
SQL_COLS = ["AC Type", "Actual Arrival Color", "Actual Arrival Date", "Actual Dep Color", "Actual Departure Date",
            "crtarpcod", "Destination", "fltcaridr", "I", "M", "Origin", "Schedule Departure Date/Time",
            "Scheduled Arrival Date", "Scheduled Departure Date", "Scheduled Departure Date (Exclude Nulls)",
            "Scheduled Departure Time", "TLM", "trkorddat", "Truck", "Vendor Code", "Vendor Name",
            "FLIGHT_LEG_SEQ_NBR", "fltseqnum", "ITIN_FLIGHT_LEG_SEQ_NBR", "MKT_FLIGHT_NBR", "Placeholder"]


def tableau_rows(date, truck, o, d, sdep, sarr, tlm, adep, aarr):
    """Mimic the Tableau download: 3 rows per TLM, actual only on its own pane."""
    base = [date, truck, "111", "ACME TRUCKING", o, d, "DRYD", sdep, sarr, tlm]
    return [base + [adep, ""], base + ["", ""], base + ["", aarr]]


def sample() -> pd.DataFrame:
    rows = []
    rows += tableau_rows("10/17/2026", "7001", "JFK", "BOS", "10/17/2026 8:00", "10/17/2026 14:00", "T1",
                         "10/17/2026 8:30", "10/17/2026 14:30")                       # on time both
    rows += tableau_rows("10/17/2026", "7002", "JFK", "BOS", "10/17/2026 8:00", "10/17/2026 14:00", "T2",
                         "10/17/2026 9:00", "10/17/2026 15:30")                       # 60 min dep late, arr late
    rows += tableau_rows("10/17/2026", "7003", "ORD", "DFW", "10/17/2026 20:00", "10/18/2026 13:05", "T3",
                         "10/17/2026 20:59", "11/18/2026 13:00")                      # dep 59 = on time, arr bad date
    rows += tableau_rows("10/16/2026", "7004", "ORD", "DFW", "10/16/2026 20:00", "10/17/2026 13:05", "T4",
                         "", "")                                                      # no actual (past)
    rows += tableau_rows("10/18/2026", "7005", "MIA", "ATL", "10/18/2026 23:00", "10/19/2026 9:00", "T5",
                         "", "")                                                      # pending (today)
    return pd.DataFrame(rows, columns=COLS)


def sql_row(tlm, truck, o, d, sdep, sarr, adep, aarr, i="10/16/2026 3:00:00 AM", m=None, order="10/15/2026 2:00:00 PM",
            dep_color="GREEN", arr_color="GREEN"):
    sd = pd.Timestamp(sdep)
    f = lambda x: "" if not x else pd.Timestamp(x).strftime("%m/%d/%Y %I:%M:%S %p")  # noqa: E731
    v = {"AC Type": "DRYD", "Actual Arrival Color": arr_color, "Actual Arrival Date": f(aarr),
         "Actual Dep Color": dep_color, "Actual Departure Date": f(adep), "crtarpcod": o, "Destination": d,
         "fltcaridr": "1001", "I": i, "M": m or i, "Origin": o, "Schedule Departure Date/Time": f(sdep),
         "Scheduled Arrival Date": f(sarr), "Scheduled Departure Date": f"{sd:%m/%d/%Y}",
         "Scheduled Departure Date (Exclude Nulls)": f"{sd:%m/%d/%Y}",
         "Scheduled Departure Time": f"12/30/1899 {sd:%I:%M:%S %p}", "TLM": tlm, "trkorddat": order,
         "Truck": str(truck), "Vendor Code": "77848", "Vendor Name": "VENDOR A", "FLIGHT_LEG_SEQ_NBR": "1",
         "fltseqnum": "451", "ITIN_FLIGHT_LEG_SEQ_NBR": "451", "MKT_FLIGHT_NBR": str(truck), "Placeholder": "0"}
    return [v[c] for c in SQL_COLS]


def sql_sample() -> pd.DataFrame:
    rows = [
        # Friday departure scheduled to arrive Monday; arrived Saturday -> On time, not "Check date"
        sql_row("S1", 7849, "DFW", "IAH", "2026-10-16 19:00", "2026-10-19 10:30", "2026-10-16 19:10", "2026-10-17 09:24"),
        # arrival keyed a month before the truck left -> Check date
        sql_row("S2", 8020, "FLR", "MXP", "2026-10-16 20:00", "2026-10-17 10:00", "2026-10-17 01:00", "2026-09-17 05:00"),
        # departure keyed 24 h early -> Check date
        sql_row("S3", 7970, "IAD", "PHL", "2026-10-16 18:00", "2026-10-17 12:00", "2026-10-15 18:00", "2026-10-16 16:30"),
        # LHR truck at 10:00 UK time, as-of 06:00 Chicago (= 12:00 UK): not departed, overdue in UK time
        sql_row("S4", 7066, "LHR", "AMS", "2026-10-19 10:00", "2026-10-19 20:00", None, None),
        # LAX truck at 03:30 Pacific, as-of 06:00 Chicago (= 04:00 Pacific): still Pending inside the 2 h grace
        sql_row("S5", 8485, "LAX", "LAS", "2026-10-19 03:30", "2026-10-19 09:00", None, None),
        # departed 3 h late, in transit, lane usually takes 6 h -> projected late
        sql_row("S6", 7001, "JFK", "BOS", "2026-10-19 01:00", "2026-10-19 08:00", "2026-10-19 04:00", None),
        # departed, arrival due yesterday, nothing recorded -> Arrival overdue
        sql_row("S7", 7002, "JFK", "BOS", "2026-10-18 01:00", "2026-10-18 08:00", "2026-10-18 01:20", None),
        # superseded version of S1 (I <> M) with a wrong actual - must be ignored
        sql_row("S1", 7849, "DFW", "IAH", "2026-10-16 19:00", "2026-10-19 10:30", "2026-10-16 23:59", "2026-10-17 09:24",
                i="10/15/2026 3:00:00 AM", m="10/16/2026 3:00:00 AM"),
        # S8 matched two flight legs (join has no date key): keep the one near the order date
        sql_row("S8", 7100, "PHL", "BOS", "2026-10-17 20:00", "2026-10-18 06:00", "2026-10-17 20:05", "2026-10-18 05:50",
                order="10/17/2026 9:00:00 AM"),
        sql_row("S8", 7100, "PHL", "BOS", "2026-09-26 20:00", "2026-09-27 06:00", "2026-09-26 23:00", "2026-09-27 09:00",
                order="10/17/2026 9:00:00 AM"),
    ]
    # lane history for JFK-BOS: clean 6 h trips
    for k in range(5):
        d = pd.Timestamp("2026-10-10") + pd.Timedelta(days=k)
        rows.append(sql_row(f"H{k}", 7001, "JFK", "BOS", d + pd.Timedelta(hours=1), d + pd.Timedelta(hours=8),
                            d + pd.Timedelta(hours=1), d + pd.Timedelta(hours=7)))
    return pd.DataFrame(rows, columns=SQL_COLS)


class Args:
    def __init__(self, **k):
        self.days = self.date_from = self.date_to = self.origin = self.dest = self.vendor = self.truck = None
        self.threshold, self.report_card, self.as_of, self.arr_basis = 59, "none", "2026-10-18 08:00", "system"
        self.no_history, self.history_only, self.hub = True, False, None
        self.__dict__.update(k)


class LateTrucksTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.csv = self.tmp / "Sheet 1_data.csv"
        sample().to_csv(self.csv, index=False)
        self.sql = self.tmp / "custom_sql_query.csv"
        sql_sample().to_csv(self.sql, index=False)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def run_build(self, **k):
        k.setdefault("inputs", [str(self.csv)])
        return L.build(Args(out=str(self.tmp / "out"), **k), log=lambda *_: None)

    def run_sql(self, **k):
        k.setdefault("as_of", "2026-10-19 06:30")
        return self.run_build(inputs=[str(self.sql)], **k)

    # ---- Tableau "Download > Data" layout
    def test_collapse_one_row_per_tlm(self):
        raw = L.map_columns(L.read_any(self.csv))
        df = L.collapse(raw)
        self.assertEqual(len(raw), 15)
        self.assertEqual(len(df), 5)
        t2 = df[df.tlm == "T2"].iloc[0]
        self.assertEqual(t2.act_dep, "10/17/2026 9:00")
        self.assertEqual(t2.act_arr, "10/17/2026 15:30")

    def test_statuses(self):
        df = self.run_build()["df"].set_index("tlm")
        self.assertEqual(df.loc["T1", "dep_status"], "On time")
        self.assertEqual(df.loc["T2", "dep_status"], "Late")       # 60 > 59
        self.assertEqual(df.loc["T2", "arr_status"], "Late")
        self.assertEqual(df.loc["T3", "dep_status"], "On time")    # 59 is not late
        self.assertEqual(df.loc["T3", "arr_status"], "Check date")
        self.assertEqual(df.loc["T4", "dep_status"], "No actual")
        self.assertEqual(df.loc["T5", "dep_status"], "Pending")
        self.assertEqual(int(df.loc["T2", "dep_delay"]), 60)

    def test_kpis_and_outputs(self):
        res = self.run_build()
        k = res["feed"]["kpis"]
        self.assertEqual((k["trucks"], k["dep_late"], k["arr_late"], k["no_dep"], k["check_date"]), (5, 1, 1, 1, 1))
        self.assertTrue(res["xlsx"].exists())
        self.assertTrue((self.tmp / "out" / "late_trucks_feed.js").exists())
        import openpyxl
        wb = openpyxl.load_workbook(res["xlsx"])
        self.assertEqual(wb.sheetnames, ["Dashboard", "Worklist", "Report Card", "Summary", "Trends", "Data", "Lanes",
                                         "Vendors", "Stations", "Settings", "About"])
        ws = wb["Data"]
        hdr = {c.value: c.column for c in ws[3]}
        self.assertTrue(str(ws.cell(4, hdr["Dep Status"]).value).startswith("=IF(Data[[#This Row],[ATD]]"))
        cached = openpyxl.load_workbook(res["xlsx"], data_only=True)["Data"]
        status = {cached.cell(r, hdr["TLM"]).value: cached.cell(r, hdr["Dep Status"]).value
                  for r in range(4, ws.max_row + 1)}
        self.assertEqual(status["T2"], "Late")
        self.assertEqual(status["T1"], "On time")
        rc = wb["Report Card"]
        self.assertEqual([c.value for c in rc[4]][:23], L.RC_COLUMNS)
        rcc = openpyxl.load_workbook(res["xlsx"], data_only=True)["Report Card"]
        trucks = {str(rcc.cell(r, 5).value) for r in range(5, rc.max_row + 1)}
        self.assertEqual(trucks, {"7002", "7003", "7004", "7005"})
        s = openpyxl.load_workbook(res["xlsx"], data_only=True)["Summary"]
        self.assertEqual(s["B6"].value, 5)                                  # trucks in window
        self.assertTrue(str(wb["Summary"]["B6"].value).startswith("=COUNTIFS("))

    def test_worklist_and_carry_forward(self):
        import openpyxl
        res = self.run_build()
        wl = openpyxl.load_workbook(res["xlsx"])
        ws = wl["Worklist"]
        hdr = {c.value: c.column for c in ws[4]}
        tlms = [ws.cell(r, hdr["TLM"]).value for r in range(5, ws.max_row + 1)]
        self.assertEqual(set(tlms), {"T2", "T3", "T4"})                       # late, check date, no ATD
        cached = openpyxl.load_workbook(res["xlsx"], data_only=True)["Worklist"]
        issue = {cached.cell(r, hdr["TLM"]).value: cached.cell(r, hdr["Issue"]).value for r in range(5, ws.max_row + 1)}
        self.assertEqual(issue["T4"], "No ATD - possible no-show")
        self.assertEqual(cached.cell(5, hdr["TLM"]).value, "T4")             # priority 1 first
        r = tlms.index("T2") + 5
        ws.cell(r, hdr["Status"]).value = "Emailed vendor"
        ws.cell(r, hdr["Delay Reason"]).value = "MECH"
        ws.cell(r, hdr["Notes"]).value = "flat tire"
        v = wl["Vendors"]
        v.cell(5, 3).value = "dispatch@acme.example"                          # vendor email typed once
        wl.save(res["xlsx"])
        res2 = self.run_build(as_of="2026-10-18 09:00")
        self.assertEqual(res2["notes"]["T2"]["Notes"], "flat tire")
        c2 = openpyxl.load_workbook(res2["xlsx"], data_only=True)
        w2 = c2["Worklist"]
        row = {w2.cell(r, hdr["TLM"]).value: r for r in range(5, w2.max_row + 1)}["T2"]
        self.assertEqual(w2.cell(row, hdr["Status"]).value, "Emailed vendor")
        self.assertEqual(w2.cell(row, hdr["Fault"]).value, "VN")
        rc = c2["Report Card"]
        rc_row = [r for r in range(5, rc.max_row + 1) if rc.cell(r, 25).value == "T2"][0]
        self.assertEqual(rc.cell(rc_row, 19).value, "MECH")                    # Reason linked from the Worklist
        self.assertEqual(rc.cell(rc_row, 21).value, "flat tire")              # Summary <- Notes
        self.assertEqual(c2["Vendors"].cell(5, 3).value, "dispatch@acme.example")
        self.assertEqual(w2.cell(row, hdr["Email vendor"]).value, "Email vendor")
        f2 = openpyxl.load_workbook(res2["xlsx"])["Worklist"].cell(row, hdr["Email vendor"]).value
        self.assertIn("mailto:", f2)

    def test_filters(self):
        self.assertEqual(len(self.run_build(origin="ORD")["df"]), 2)
        self.assertEqual(len(self.run_build(days=1)["df"]), 1)
        self.assertEqual(len(self.run_build(date_from="2026-10-17", date_to="2026-10-17")["df"]), 3)

    def test_rc_time(self):
        d = pd.Timestamp("2026-10-17")
        self.assertEqual(L.rc_time(pd.Timestamp("2026-10-17 20:30"), d), "2030")
        self.assertEqual(L.rc_time(pd.Timestamp("2026-10-18 13:05"), d), "1305+1")
        self.assertEqual(L.rc_time(pd.NaT, d), "")

    def test_team_schedule_basis(self):
        df = L.score(L.prepare(L.collapse(L.map_columns(L.read_any(self.csv))), datetime(2026, 10, 18, 8)), 59)
        sched = {("7001", "JFK", "BOS"): (480, 12 * 60)}          # team ETA 12:00 vs system 14:00
        out = L.apply_team_schedule(df, sched, "team", 59).set_index("tlm")
        self.assertEqual(out.loc["T1", "arr_status"], "Late")    # 14:30 vs 12:00
        self.assertEqual(out.loc["T2", "arr_status"], "Late")    # not on schedule -> system basis kept
        self.assertEqual(L._hhmm_offset("1305+1"), 13 * 60 + 5 + 1440)
        self.assertEqual(L._hhmm_offset(400), 240)

    def test_minutes_like_tableau_datediff(self):
        a = pd.Series(pd.to_datetime(["2026-10-17 08:59:59", "2026-10-17 09:00:00"]))
        b = pd.Series(pd.to_datetime(["2026-10-17 08:00:00", "2026-10-17 08:00:30"]))
        self.assertEqual(list(L._minutes(a, b)), [59, 60])

    # ---- custom-SQL data source layout
    def test_sql_layout_rules(self):
        df = self.run_sql()["df"].set_index("tlm")
        self.assertEqual(len(df.loc[["S1"]]), 1)                       # superseded version dropped
        self.assertEqual(df.loc["S1", "act_dep"], pd.Timestamp("2026-10-16 19:10"))
        self.assertEqual(df.loc["S1", "arr_status"], "On time")        # weekend ETA, arrived early
        self.assertEqual(df.loc["S2", "arr_status"], "Check date")     # arrival before departure
        self.assertIn("Arrival before departure", df.loc["S2", "issue"])
        self.assertEqual(df.loc["S3", "dep_status"], "Check date")     # 24 h early
        self.assertEqual(df.loc["S8", "sched_date"], pd.Timestamp("2026-10-17"))  # leg nearest the order date
        self.assertEqual(df.loc["S8", "dep_status"], "On time")

    def test_station_local_time_and_actions(self):
        df = self.run_sql()["df"].set_index("tlm")
        self.assertEqual(df.loc["S4", "dep_status"], "No actual")      # 12:00 UK > 10:00 + 2 h grace
        self.assertEqual(df.loc["S4", "action"], "not_departed")
        self.assertEqual(df.loc["S5", "dep_status"], "Pending")        # 04:00 PT < 03:30 + 2 h grace
        self.assertEqual(df.loc["S6", "action"], "proj_late")          # 04:00 + 6 h lane median = 10:00 vs 08:00
        self.assertEqual(df.loc["S6", "proj_arr"], pd.Timestamp("2026-10-19 10:00"))
        self.assertEqual(int(df.loc["S6", "proj_arr_delay"]), 120)
        self.assertEqual(df.loc["S7", "action"], "arr_overdue")
        self.assertTrue(pd.isna(df.loc["H4", "action"]))

    def test_tableau_color_check(self):
        res = self.run_sql()
        tc = L.tableau_check(res["df_all"])
        self.assertEqual(tc["dep"]["check_date_tableau_colored"], 1)  # S3

    # ---- history store
    def test_history_store_merges_and_rereads(self):
        hist = self.tmp / "hist"
        self.run_sql(no_history=False, history_dir=str(hist), history_only=True)
        h = pd.read_csv(hist / L.HISTORY_FILE, dtype=str)
        self.assertEqual(len(h), 13)
        self.assertEqual(h["TLM"].duplicated().sum(), 0)
        self.assertTrue((hist / "late_trucks_daily.csv").exists())
        self.assertTrue((hist / "Trucks History Summary.xlsx").exists())
        # newer export: S4 has now departed -> replaces the stored row, nothing duplicated
        time.sleep(0.01)
        newer = self.tmp / "custom_sql_query_2.csv"
        pd.DataFrame([sql_row("S4", 7066, "LHR", "AMS", "2026-10-19 10:00", "2026-10-19 20:00",
                              "2026-10-19 12:30", None)], columns=SQL_COLS).to_csv(newer, index=False)
        os.utime(newer, (time.time() + 5, time.time() + 5))
        res = self.run_build(inputs=[str(newer)], no_history=False, history_dir=str(hist), as_of="2026-10-19 08:00")
        h = pd.read_csv(hist / L.HISTORY_FILE, dtype=str).set_index("TLM")
        self.assertEqual(len(h), 13)
        self.assertEqual(h.loc["S4", "Dep Status"], "Late")
        self.assertEqual(h.loc["S1", "Arr Status"], "On time")         # untouched rows survive the round trip
        self.assertEqual(h.loc["S1", "Source File"], "custom_sql_query.csv")
        self.assertEqual(res["df"].set_index("tlm").loc["S4", "dep_delay"], 150)


if __name__ == "__main__":
    unittest.main()
