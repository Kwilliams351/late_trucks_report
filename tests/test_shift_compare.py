import shutil
import tempfile
import unittest
from pathlib import Path

import openpyxl
import pandas as pd

import late_trucks as L
import shift_compare as S
from report_card_pending import CARD_STATUSES, card_status
from test_late_trucks import Args, sql_sample


class ShiftCompareTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.csv = self.tmp / "input.csv"
        sql_sample().to_csv(self.csv, index=False)

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def build(self, stamp="2026-10-19 06:30", **kwargs):
        return L.build(Args(inputs=[str(self.csv)], out=str(self.tmp / "Output"), as_of=stamp, **kwargs),
                       log=lambda *_: None)

    def test_six_statuses_and_priority(self):
        cases = [
            ("Late", "Late", "2026-10-19 05:00", "2026-10-19 10:00", "Normal", 100, True),
            ("Late", "No actual", "2026-10-19 05:00", None, "Normal", 100, True),
            ("On time", "Pending", "2026-10-19 05:00", None, "Normal", 60, False),
            ("On time", "Pending", "2026-10-19 05:00", None, "Normal", 59, False),
            ("No actual", "No actual", None, None, "Normal", None, False),
            ("Late", "Check date", "2026-10-19 05:00", "2026-09-19 10:00", "Normal", None, True),
            ("Late", "Not tracked", "2026-10-19 05:00", None, "None", None, True),
            ("On time", "On time", "2026-10-19 05:00", "2026-10-19 10:00", "Normal", None, False),
            ("Pending", "Pending", None, None, "Normal", None, False),
        ]
        df = pd.DataFrame(cases, columns=["dep_status", "arr_status", "act_dep", "act_arr",
                                          "arr_rule", "proj_arr_delay", "late"])
        df["act_dep"] = pd.to_datetime(df["act_dep"])
        df["act_arr"] = pd.to_datetime(df["act_arr"])
        self.assertEqual(list(card_status(df, 59)), [*CARD_STATUSES, "Ready to log", "", "Pending - no departure"])
        self.assertEqual(card_status(df, 60).iloc[2], "Pending - in transit")

    def test_report_rows_formulas_and_projections(self):
        res = self.build()
        wb = openpyxl.load_workbook(res["xlsx"])
        cached = openpyxl.load_workbook(res["xlsx"], data_only=True)
        try:
            rc = cached["Report Card"]
            headers = {cell.value: cell.column for cell in cached["Data"][3]}
            actual = {rc.cell(r, 25).value: rc.cell(r, 24).value for r in range(5, rc.max_row + 1)}
            expected = card_status(res["df_data"], res["settings"]["threshold"])
            expected = dict(zip(res["df_data"].tlm, expected))
            for tlm, status in actual.items():
                self.assertEqual(status, expected[tlm])
            self.assertEqual(actual["S6"], "Pending - projected late")
            self.assertEqual(actual["S7"], "Pending - arrival overdue")
            self.assertEqual(actual["S2"], "Check record")
            self.assertEqual(actual["S5"], "Pending - no departure")
            ordered = list(actual.values())
            self.assertEqual(ordered, sorted(ordered, key=CARD_STATUSES.index))
            self.assertEqual(rc["X4"].value, "Card status")
            self.assertEqual(rc["AF4"].value, "Proj ATA")
            self.assertEqual(rc["AG4"].value, "Proj Arr Delay (min)")
            for r in range(5, rc.max_row + 1):
                if rc.cell(r, 25).value == "S5":
                    self.assertEqual(rc.cell(r, 28).value, "Not late yet")
            status_col = headers["Card status"]
            self.assertIn("Threshold", wb["Data"].cell(4, status_col).value)
            self.assertIn("INDEX(Data[Card status]", wb["Report Card"]["X5"].value)
            self.assertIn("COUNTIF", wb["Report Card"]["A3"].value)
            for ws in wb:
                for row in ws:
                    for cell in row:
                        self.assertNotIn("#REF!", str(cell.value))
            for ws in cached:
                for row in ws:
                    for cell in row:
                        self.assertNotEqual(cell.data_type, "e")
        finally:
            wb.close()
            cached.close()

    def test_snapshot_round_trip_and_first_run(self):
        res = self.build()
        previous = S.read_snapshot(res["snapshot"])
        workbook = S.read_snapshot(res["xlsx"])
        pd.testing.assert_frame_equal(previous.set_index("TLM"), workbook.set_index("TLM"), check_like=True)
        self.assertIsNone(res["shift_comparison"])
        wb = openpyxl.load_workbook(res["xlsx"])
        self.assertNotIn("Shift Variance", wb.sheetnames)
        wb.close()
        again = self.build()
        self.assertIsNone(again["shift_comparison"])
        self.assertNotEqual(res["snapshot"], again["snapshot"])
        self.assertIsNone(S.select_prior(self.tmp / "snapshots", res["as_of"]))

    def test_clock_only_and_real_updates(self):
        first = self.build()
        later = self.build("2026-10-19 14:30")
        comp = later["shift_comparison"]
        self.assertIsNotNone(comp)
        self.assertGreater(comp["transitions"]["Clock only"], 0)
        self.assertGreater(comp["transitions"]["Now overdue"], 0)
        wb = openpyxl.load_workbook(later["xlsx"])
        self.assertIn("Shift Variance", wb.sheetnames)
        self.assertGreater(len(wb["Dashboard"]._charts), 0)
        wb.close()
        before = S.read_snapshot(first["snapshot"])
        after = before.copy()
        after.attrs["as_of"] += pd.Timedelta(hours=1)
        row = after.index[after.TLM == "S6"][0]
        after.loc[row, ["ATA", "Arr Status", "Card status"]] = [
            "2026-10-19 11:00", "Late", "Ready to log"]
        result = S.compare_snapshots(before, after)
        self.assertEqual(result["transitions"]["Pending -> ready to log"], 1)
        self.assertEqual(result["transitions"]["Became late"], 0)  # S6 already departed late
        self.assertEqual(result["transitions"]["Clock only"], 0)
        after.loc[row, ["STA", "Dep Status", "Arr Status", "Card status"]] = [
            "2026-10-19 12:00", "On time", "On time", ""]
        regained = S.compare_snapshots(before, after)
        self.assertEqual(regained["transitions"]["No longer late (time regained)"], 1)
        after.loc[row, "Arr Status"] = "Check date"
        self.assertEqual(S.compare_snapshots(before, after)["transitions"]["No longer late (time regained)"], 0)

    def test_prior_selection_and_standalone(self):
        first = self.build("2026-10-18 06:30")
        second = self.build("2026-10-19 06:00")
        now = self.build("2026-10-19 06:30", vs="prior-day")
        self.assertEqual(now["shift_comparison"]["prior_as_of"], pd.Timestamp(first["as_of"]))
        prior = S.select_prior(self.tmp / "snapshots", now["as_of"])
        self.assertEqual(prior.attrs["as_of"], pd.Timestamp(second["as_of"]))
        with self.assertRaises(ValueError):
            S.select_prior(self.tmp / "snapshots", first["as_of"], str(second["xlsx"]))
        target = self.tmp / "variance.xlsx"
        S.main(["compare", str(now["xlsx"]), "--vs", str(first["xlsx"]), "--out", str(target)])
        wb = openpyxl.load_workbook(target)
        self.assertEqual(wb.sheetnames, ["Shift Variance"])
        wb.close()

    def test_added_removed_and_empty_snapshot(self):
        res = self.build()
        old = S.read_snapshot(res["snapshot"])
        new = old.iloc[1:].copy()
        new.attrs["as_of"] += pd.Timedelta(hours=1)
        result = S.compare_snapshots(old, new)
        self.assertEqual(result["changes"][0]["Changed"], "Removed")
        self.assertEqual(sum(result["transitions"].values()), 0)
        empty = new.iloc[:0].copy()
        path = S.save_snapshot(empty, self.tmp / "empty")
        self.assertTrue(S.read_snapshot(path).empty)
        self.assertEqual(S.compare_snapshots(old, empty)["current_totals"]["Trucks"], 0)

    def test_invalid_cached_values_and_duplicate_keys_rejected(self):
        res = self.build()
        original = S.read_snapshot(res["snapshot"])
        for field, value in (("TLM", ""), ("Dep Status", ""), ("Arr Status", ""),
                             ("Card status", "unknown"), ("In Window", "")):
            invalid = original.copy()
            invalid.loc[0, field] = value
            with self.assertRaises(ValueError):
                S._validate(invalid, res["as_of"])
        duplicate = pd.concat([original, original.iloc[:1]], ignore_index=True)
        with self.assertRaises(ValueError):
            S._validate(duplicate, res["as_of"])

    def test_became_late_and_rules_are_not_clock_only(self):
        res = self.build()
        previous = S.read_snapshot(res["snapshot"])
        current = previous.copy()
        current.attrs["as_of"] += pd.Timedelta(hours=1)
        row = current.index[current.TLM == "S5"][0]
        current.loc[row, ["ATD", "ATA", "Dep Status", "Arr Status", "Card status"]] = [
            "2026-10-19 05:30", "2026-10-19 11:00", "Late", "Late", "Ready to log"]
        result = S.compare_snapshots(previous, current)
        self.assertEqual(result["transitions"]["Became late"], 1)
        self.assertEqual(result["transitions"]["Pending -> ready to log"], 1)
        self.assertEqual(result["transitions"]["Clock only"], 0)
        current = previous.copy()
        current.attrs["as_of"] += pd.Timedelta(hours=1)
        row = current.index[current.TLM == "S6"][0]
        current.loc[row, ["Card status", "Threshold"]] = ["Pending - in transit", "180"]
        self.assertEqual(S.compare_snapshots(previous, current)["transitions"]["Clock only"], 0)


if __name__ == "__main__":
    unittest.main()
