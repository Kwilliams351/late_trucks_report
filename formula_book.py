"""Formula-linked workbooks for the late trucks report.

Every number in these workbooks traces back to one Data tab through Excel formulas:

    Tableau export --(late_trucks.py: clean, one row per TLM)--> Data tab (blue headers = values from the export)
        Data tab formula columns (green headers): delays, statuses, flags, issue          [row level]
        Summary / Trends / Lanes / Stations: COUNTIFS / SUMIFS / MEDIAN over Data          [tables]
        Dashboard: KPI tiles and charts that point at Summary / Trends / Worklist          [visuals]
        Worklist / Report Card: INDEX into Data by row number + the coordinator's inputs   [work]

Rules (threshold, grace, check-date limits, as-of time, report window) live on Settings as named cells, so changing
one there re-scores the whole workbook. Python writes the same results as cached values so the file reads
correctly even before Excel recalculates (e.g. in Protected View).
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from report_card_pending import CARD_STATUSES, card_status, card_status_formula

EXCEL_EPOCH = pd.Timestamp("1899-12-30")
RED, GREEN, AMBER, GREY = "#E15759", "#59A14F", "#B07A00", "#8C8C8C"
BLUE, FORM, INPUT = "#4D77A4", "#2E7D5B", "#FFF2CC"
NAVY = "#1F3864"

ISSUES = [  # (issue, priority, meaning)
    ("No ATD - possible no-show", 1, "Scheduled departure passed (plus grace) and no actual departure recorded"),
    ("No ATA - arrival overdue", 1, "Departed, scheduled arrival passed (plus grace), no actual arrival recorded"),
    ("Late arrival", 2, "Actual arrival more than the threshold after the scheduled arrival"),
    ("Projected late", 2, "In transit; actual departure + lane median transit lands after STA + threshold"),
    ("Late departure", 3, "Actual departure more than the threshold after the scheduled departure"),
    ("Fix record - date keyed wrong", 4, "An actual was keyed on the wrong day/month (see Check date)"),
]
STEPS = [  # (key, what to do) - from the Trucking Training Guide (Tableau, June 2023) and No Shows (Jan 2025)
    ("NOATD", "Possible no-show. TRK001: confirm there is no ATD. Search the trucking team inbox (truck number, "
              "AWB, date). No status: email the vendor for ETA + reason (vendor in TRK001). If the truck did not run, "
              "follow No Shows: CAP142 (-2 to +1 days, Truck, status New) > OPR352 check AWBs not tendered > CAP018 "
              "cancel trip segments (not offline stations) > scribble pad 'no show/date/initials/trucking team'."),
    ("NOATA", "TRK001 (or FLT002) for ATA. Search the trucking team inbox. No status: email the vendor for ETA + "
              "reason and update the station. To arrive it: Station = {DEST}, Role = CTRKCRD, TRK001 > Reopen > "
              "enter actual arrival date/time > Save > Approve."),
    ("NOATA_OFFLINE", "{DEST} is an OFFLINE station (no iCargo). OPR352 (AWB Enquiry) > three dots > FSU messages > "
                      "RCF for the ATA. No FSU: email the station for the ATA. Then update TRK001 with the ATA and "
                      "arrive the shipment in LTE003."),
    ("NOATA_MAIL", "Mail truck: arrival time can be entered as the STA. TRK001 > Reopen > ATA = STA > Save > Approve "
                   "(Station = {DEST}, Role = CTRKCRD)."),
    ("LATE", "Confirm STD/STA/ATD/ATA in TRK001. Check the trucking team inbox, then email the vendor for the "
             "cause of delay; update the station. Log on the Report Card (Reason / Summary)."),
    ("PROJ", "In transit and running late: email the vendor for an ETA and warn {DEST}. Re-check TRK001 / FLT002 "
             "after the STA."),
    ("FIX", "An actual looks keyed on the wrong date. Verify in TRK001 / FLT002, then correct it: TRK001 > Reopen > "
            "fix ATD/ATA > Save > Approve."),
]
STATUSES = ["Not started", "Checked iCargo", "Emailed vendor", "Waiting on vendor", "Emailed station",
            "Updated in iCargo", "Logged on Report Card", "Closed - no action"]
DONE = ["Updated in iCargo", "Logged on Report Card", "Closed - no action"]
REASONS = [("MECH", "VN"), ("DRIVER", "VN"), ("VENDOR", "VN"), ("CUSTOMS", "CB"), ("AA OPS", "AA"),
           ("AA SALES", "AA"), ("WX", "WX")]
TYPE_ROWS = [("XFR", "Transfer", "None", "Transfer trucks never have arrival times (training guide)"),
             ("RFS", "RFS", "None", "No arrival on the flight record"),
             ("MAL", "Mail", "STA", "Mail truck arrival is entered as the STA; a missing one is still chased")]

# Offline stations (no iCargo) - Trucking Training Guide / No Shows, Jan 2025
OFFLINE = {
    "Australia": "BNE MEL", "Belgium": "BRU", "Brazil": "CNF CWB POA VCP", "Denmark": "BLL CPH",
    "Dominican Republic": "SDQ", "France": "BOD LIL LYS MLH MPL MRS NCE NTE ORY RNS SXB TLS",
    "Germany": "BRE CGN DRS DTM DUS FMO HAJ HAM LEJ LUX NUE STR TXL", "Ireland": "ORK SNN",
    "Italy": "AOI BLQ FLR GOA NAP MIL TRN VRN", "Japan": "KIX NGO", "Mexico": "MTY", "Netherlands": "EIN MST RTM",
    "Norway": "OSL", "Portugal": "OPO", "Spain": "AGP ALC BIO SVQ VIT VLC ZAZ", "Sweden": "ARN GOT MMA",
    "Switzerland": "BRN BSL GVA", "UK": "ABZ BFS BHX BRS CWL EDI EMA GLA LBA NCL", "USA": "GSO CVG LRD ROC",
}
OFFLINE_NOTES = {
    "BLL": "Outbound trucking service only", "SDQ": "2 locations in SDQ: 1 cargo terminal, 1 truck terminal",
    "TXL": "Trucking only until 6/7/2019; then flights + trucks (summer-only flights)",
    "SNN": "Seasonal flights; otherwise trucking only",
    "BLQ": "Trucking only until 6/7/2019; then flights + trucks (summer-only flights)",
    "MTY": "Outbound MTY truck only, inbound AE flights only", "OSL": "Outbound trucking service only",
    "OPO": "Outbound trucking service only", "ARN": "Outbound trucking service only",
    "GOT": "Outbound trucking service only", "BHX": "Seasonal flights; otherwise trucking only",
    "BRS": "Inbound trucking only", "EDI": "Seasonal flights; otherwise trucking only",
    "GLA": "Seasonal flights; otherwise trucking only", "GSO": "Drop station - handler Forward Air vendor",
    "CVG": "AE flights only / PPS at BSO; cargo at trucking terminal / handler PACE Trucking",
    "LRD": "Outbound cargo is truck only / handler Campbells Delivery Service; inbound cargo/PPS on AE flights, "
           "picked up at BSO",
    "ROC": "Inbound trucking only / handler Jetstream Ground Services; PPS avail at BSO",
}
OFFLINE_REGION = {code: region for region, codes in OFFLINE.items() for code in codes.split()}


# ------------------------------------------------------------------------------------------------ helpers
def _speed_up_xlsxwriter():
    """xlsxwriter runs ~30 regexes on every formula; the Data tab repeats the same formula on every row, so memoize."""
    from xlsxwriter.worksheet import Worksheet
    if getattr(Worksheet._prepare_formula, "_memo", False):
        return
    orig = Worksheet._prepare_formula
    cache = {}

    def fast(self, formula, expand_future_functions=False):
        k = (formula, expand_future_functions, self.use_future_functions)
        v = cache.get(k)
        if v is None:
            v = cache[k] = orig(self, formula, expand_future_functions)
            if len(cache) > 50000:
                cache.clear()
        return v

    fast._memo = True
    Worksheet._prepare_formula = fast


def serial(ts):
    """Excel serial number for a timestamp (cached values of date formulas)."""
    if ts is None or pd.isna(ts):
        return ""
    return (pd.Timestamp(ts) - EXCEL_EPOCH) / pd.Timedelta(days=1)


def col_letter(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def R(col: str) -> str:
    """This-row reference inside the Data table."""
    return f"Data[[#This Row],[{col}]]"


def _v(x):
    """Cached value for a formula cell (never None - xlsxwriter would write the text 'None')."""
    if x is None:
        return ""
    if isinstance(x, (pd.Timestamp, datetime)):
        return serial(x)
    if isinstance(x, (bool, np.bool_)):
        return int(x)
    try:
        if pd.isna(x):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(x, (np.integer,)):
        return int(x)
    if isinstance(x, (np.floating,)):
        return float(x)
    return x


def ddmmm(ts) -> str:
    return "" if ts is None or pd.isna(ts) else pd.Timestamp(ts).strftime("%d%b").upper()


def rc_time(ts, base) -> str:
    if ts is None or pd.isna(ts) or pd.isna(base):
        return ""
    d = (pd.Timestamp(ts).normalize() - pd.Timestamp(base).normalize()).days
    return f"{pd.Timestamp(ts):%H%M}" + (f"+{d}" if d > 0 else (f"-{-d}" if d < 0 else ""))


class Fmt:
    def __init__(self, wb):
        self.wb = wb
        F = self.F
        self.title = F(bold=True, font_size=16, font_color="white", bg_color=NAVY)
        self.title_r = F(bold=True, font_color="white", bg_color=NAVY, align="right")
        self.sub = F(font_color="#666666", italic=True)
        self.h2 = F(bold=True, font_size=12, font_color=NAVY)
        self.hv = F(bold=True, font_color="white", bg_color=BLUE, text_wrap=True, valign="bottom", border=1,
                    border_color="#D0D0D0")
        self.hf = F(bold=True, font_color="white", bg_color=FORM, text_wrap=True, valign="bottom", border=1,
                    border_color="#D0D0D0")
        self.hi = F(bold=True, font_color="#444444", bg_color="#FFD966", text_wrap=True, valign="bottom", border=1,
                    border_color="#D0D0D0")
        self.txt = F()
        self.wrap = F(text_wrap=True, valign="top")
        self.num = F(num_format="0")
        self.num1 = F(num_format="0.0")
        self.int_ = F(num_format="#,##0")
        self.pct = F(num_format="0.0%")
        self.date = F(num_format="m/d/yyyy")
        self.ddd = F(num_format="ddd m/d")
        self.dtm = F(num_format="m/d/yy hh:mm")
        self.inp = F(bg_color=INPUT, border=1, border_color="#E0C060")
        self.inp_dt = F(bg_color=INPUT, border=1, border_color="#E0C060", num_format="m/d/yy hh:mm")
        self.inp_wrap = F(bg_color=INPUT, border=1, border_color="#E0C060", text_wrap=True, valign="top")
        self.red = F(font_color=RED, bold=True)
        self.green = F(font_color=GREEN)
        self.amber = F(font_color=AMBER, italic=True)
        self.grey = F(font_color=GREY, italic=True)
        self.greyfill = F(font_color="#999999", bg_color="#F2F2F2")
        self.tile_h = F(bold=True, font_color="white", bg_color="#595959", align="center", valign="vcenter",
                        text_wrap=True)
        self.tile_v = F(bold=True, font_size=20, align="center", valign="vcenter", num_format="#,##0",
                        font_color=NAVY, border=1, border_color="#D9D9D9")
        self.tile_p = F(bold=True, font_size=20, align="center", valign="vcenter", num_format="0.0%",
                        font_color=NAVY, border=1, border_color="#D9D9D9")
        self.tile_r = F(bold=True, font_size=20, align="center", valign="vcenter", num_format="#,##0",
                        font_color=RED, border=1, border_color="#D9D9D9")

    def F(self, **k):
        return self.wb.add_format({"font_name": "Arial", "font_size": 10, **k})


def head(ws, f, name, ncols, as_of_txt, note=None):
    ws.merge_range(0, 0, 0, max(ncols - 3, 1), name, f.title)
    ws.merge_range(0, max(ncols - 2, 2), 0, max(ncols - 1, 3), f"As of {as_of_txt}", f.title_r)
    if note:
        ws.write(1, 0, note, f.sub)
    ws.set_row(0, 26)


def status_colors(ws, f, rng):
    for text, fmt in (("Late", f.red), ("On time", f.green), ("Check date", f.amber), ("No actual", f.grey),
                      ("Not tracked", f.grey)):
        ws.conditional_format(rng, {"type": "cell", "criteria": "==", "value": f'"{text}"', "format": fmt})


# ------------------------------------------------------------------------------------------ data columns
def data_columns(df: pd.DataFrame, s: dict) -> list:
    """(header, kind, width, format key, formula or None, python values) for the Data table.
    kind 'v' = value from the export, 'f' = Excel formula (python values are the cached results)."""
    thr = s["threshold"]
    week = df["sched_date"] - pd.to_timedelta(df["sched_date"].dt.weekday, unit="D")
    o_off = (df["now_o"] - pd.Timestamp(s["as_of"])).dt.total_seconds() / 3600
    d_off = (df["now_d"] - pd.Timestamp(s["as_of"])).dt.total_seconds() / 3600
    dep_ok = df["dep_status"].isin(["Late", "On time"])
    planned = (df["sched_arr"] - df["sched_dep"]).dt.total_seconds() / 60
    lane_med = df["lane"].map(s["lane_median"])
    no_atd = df["dep_status"] == "No actual"
    cols = [
        ("TLM", "v", 12, "txt", None, df["tlm"]),
        ("Date", "v", 10, "date", None, df["sched_date"]),
        ("Week", "f", 10, "date", f"={R('Date')}-WEEKDAY({R('Date')},3)", week),
        ("Truck", "v", 9, "txt", None, df["truck"]),
        ("Vendor Code", "v", 9, "txt", None, df["vendor_code"]),
        ("Vendor", "v", 26, "txt", None, df["vendor"]),
        ("Origin", "v", 7, "txt", None, df["origin"]),
        ("Dest", "v", 7, "txt", None, df["dest"]),
        ("Lane", "f", 10, "txt", f'={R("Origin")}&"-"&{R("Dest")}', df["lane"]),
        ("AC Type", "v", 7, "txt", None, df["ac_type"]),
        ("Truck Type", "f", 9, "txt", f'=IFERROR(VLOOKUP(LEFT({R("AC Type")},3),TypeTbl,2,FALSE),"Standard")',
         df["truck_type"]),
        ("Arr Rule", "f", 7, "txt", f'=IFERROR(VLOOKUP(LEFT({R("AC Type")},3),TypeTbl,3,FALSE),"Normal")',
         df["arr_rule"]),
        ("STD", "v", 13, "dtm", None, df["sched_dep"]),
        ("ATD", "v", 13, "dtm", None, df["act_dep"]),
        ("Dep Delay (min)", "f", 8, "num", f'=IF({R("ATD")}="","",ROUND(({R("ATD")}-{R("STD")})*1440,0))',
         df["dep_delay"]),
        ("Dep Status", "f", 10, "txt",
         f'=IF({R("ATD")}="",IF({R("STD")}+Grace/1440<AsOf+{R("Origin Offset (h)")}/24,"No actual","Pending"),'
         f'IF(OR({R("Dep Delay (min)")}>LateCheck,{R("Dep Delay (min)")}<-EarlyCheck),"Check date",'
         f'IF({R("Dep Delay (min)")}>Threshold,"Late","On time")))', df["dep_status"]),
        ("STA", "v", 13, "dtm", None, df["sched_arr"]),
        ("ATA", "v", 13, "dtm", None, df["act_arr"]),
        ("Arr Delay (min)", "f", 8, "num", f'=IF(OR({R("ATA")}="",{R("STA")}=""),"",ROUND(({R("ATA")}-{R("STA")})*1440,0))',
         df["arr_delay"]),
        ("Transit (min)", "f", 8, "num",
         f'=IF({R("ATA")}="","",ROUND(({R("ATA")}-IF(OR({R("Dep Status")}="Late",{R("Dep Status")}="On time"),'
         f'{R("ATD")},{R("STD")}))*1440,0))', df["transit_min"]),
        ("Arr Status", "f", 10, "txt",
         f'=IF(OR({R("Arr Rule")}="None",AND({R("Arr Rule")}="STA",{R("ATA")}<>"")),"Not tracked",'
         f'IF({R("ATA")}="",IF({R("STA")}+Grace/1440<AsOf+{R("Dest Offset (h)")}/24,"No actual","Pending"),'
         f'IF(OR(N({R("Arr Delay (min)")})>LateCheck,{R("Transit (min)")}<MinTransit),"Check date",'
         f'IF(N({R("Arr Delay (min)")})>Threshold,"Late","On time"))))', df["arr_status"]),
        ("Origin Offset (h)", "f", 7, "num1", f'=IFERROR(VLOOKUP({R("Origin")},StationTbl,6,FALSE),0)', o_off.round(2)),
        ("Dest Offset (h)", "f", 7, "num1", f'=IFERROR(VLOOKUP({R("Dest")},StationTbl,6,FALSE),0)', d_off.round(2)),
        ("Late Dep", "f", 5, "num", f'=--({R("Dep Status")}="Late")', (df["dep_status"] == "Late").astype(int)),
        ("Late Arr", "f", 5, "num", f'=--({R("Arr Status")}="Late")', (df["arr_status"] == "Late").astype(int)),
        ("Arr Measured", "f", 6, "num", f'=--OR({R("Arr Status")}="Late",{R("Arr Status")}="On time")',
         df["arr_status"].isin(["Late", "On time"]).astype(int)),
        ("Late Any", "f", 5, "num", f'=MAX({R("Late Dep")},{R("Late Arr")})', df["late"].astype(int)),
        ("Check Date", "f", 5, "num", f'=--OR({R("Dep Status")}="Check date",{R("Arr Status")}="Check date")',
         ((df["dep_status"] == "Check date") | (df["arr_status"] == "Check date")).astype(int)),
        ("No ATD", "f", 5, "num", f'=--({R("Dep Status")}="No actual")', no_atd.astype(int)),
        ("No ATA", "f", 5, "num", f'=--AND({R("Arr Status")}="No actual",{R("Dep Status")}<>"No actual")',
         ((df["arr_status"] == "No actual") & ~no_atd).astype(int)),
        ("Lane Median (min)", "f", 8, "num1", f'=IFERROR(VLOOKUP({R("Lane")},LaneTbl,3,FALSE),"")', lane_med),
        ("Proj ATA", "f", 13, "dtm",
         f'=IFERROR(IF(AND({R("ATA")}="",OR({R("Dep Status")}="Late",{R("Dep Status")}="On time"),'
         f'{R("Arr Rule")}="Normal"),{R("ATD")}+ROUND(IF({R("Lane Median (min)")}<>"",{R("Lane Median (min)")},'
         f'({R("STA")}-{R("STD")})*1440),0)/1440,""),"")', df["proj_arr"]),
        ("Proj Arr Delay (min)", "f", 8, "num",
         f'=IF({R("Proj ATA")}="","",ROUND(({R("Proj ATA")}-{R("STA")})*1440,0))', df["proj_arr_delay"]),
        ("In Window", "f", 6, "num", f'=--AND({R("Date")}>=WinFrom,{R("Date")}<=WinTo)', df["in_window"].astype(int)),
        ("Issue", "f", 24, "txt",
         f'=IF({R("No ATD")},"No ATD - possible no-show",IF({R("No ATA")},"No ATA - arrival overdue",'
         f'IF({R("Check Date")},"Fix record - date keyed wrong",IF({R("Late Arr")},"Late arrival",'
         f'IF(AND({R("Arr Status")}="Pending",N({R("Proj Arr Delay (min)")})>Threshold),"Projected late",'
         f'IF({R("Late Dep")},"Late departure",""))))))', df["wl_issue"]),
        ("To Work", "f", 5, "num", f'=--AND({R("In Window")}=1,{R("Issue")}<>"")',
         (df["in_window"] & (df["wl_issue"] != "")).astype(int)),
        ("On Report Card", "v", 7, "txt", None, df["logged"].map({True: "Y", False: ""})),
        ("Tableau Dep Color", "v", 8, "txt", None, df.get("tab_dep_color")),
        ("Tableau Arr Color", "v", 8, "txt", None, df.get("tab_arr_color")),
        ("Card status", "f", 27, "txt", card_status_formula(R), card_status(df, s["threshold"])),
    ]
    del dep_ok, planned, thr
    return cols


def write_data(wb, f, df: pd.DataFrame, s: dict):
    ws = wb.add_worksheet("Data")
    cols = data_columns(df, s)
    n = len(df)
    fm = {"txt": f.txt, "date": f.date, "dtm": f.dtm, "num": f.num, "num1": f.num1}
    ws.write(0, 0, "Data - one row per TLM (cleaned Tableau export). Blue headers = values from the export; "
                   "green headers = Excel formulas (click a cell to see how it is derived). Settings holds the rules.",
             f.sub)
    HR = 2
    values = []
    for _, kind, _, _, _, ser in cols:
        if kind == "v" and ser is not None:
            values.append(list(ser))
        else:
            values.append([None] * n)
    rows = [[_cell(values[j][i]) for j in range(len(cols))] for i in range(n)]
    tcols = []
    for h, kind, w, fk, formula, _ in cols:
        c = {"header": h, "format": fm[fk], "header_format": f.hv if kind == "v" else f.hf}
        if kind == "f":
            c["formula"] = formula
        tcols.append(c)
    last = HR + max(n, 1)
    ws.add_table(HR, 0, last, len(cols) - 1, {"name": "Data", "columns": tcols, "data": rows or None,
                                               "style": "Table Style Light 1", "autofilter": True})
    # cached results for the formula columns
    for j, (h, kind, w, fk, formula, ser) in enumerate(cols):
        ws.set_column(j, j, w)
        if kind != "f":
            continue
        vals = list(ser) if ser is not None else [None] * n
        for i in range(n):
            ws.write_formula(HR + 1 + i, j, formula, fm[fk], _v(vals[i]))
    ws.set_row(HR, 42)
    ws.freeze_panes(HR + 1, 2)
    idx = {h: j for j, (h, *_) in enumerate(cols)}
    for h in ("Dep Status", "Arr Status"):
        c = col_letter(idx[h])
        status_colors(ws, f, f"{c}{HR + 2}:{c}{last + 1}")
    return cols


def _cell(x):
    if x is None:
        return None
    if isinstance(x, pd.Timestamp):
        return None if pd.isna(x) else x.to_pydatetime()
    try:
        if pd.isna(x):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(x, (np.integer,)):
        return int(x)
    if isinstance(x, (np.floating,)):
        return float(x)
    return x


# ------------------------------------------------------------------------------------------------ settings
def write_settings(wb, f, s: dict):
    ws = wb.add_worksheet("Settings")
    ws.set_column(0, 0, 30)
    ws.set_column(1, 1, 22)
    ws.set_column(2, 2, 16)
    ws.set_column(3, 3, 90)
    head(ws, f, "Settings - rules used by every formula", 4, s["as_of_txt"],
         "Yellow cells can be changed; the whole workbook re-scores (Data > Summary > Trends > Dashboard > Worklist).")
    params = [
        ("AsOf", "As-of time (PC clock)", s["as_of"], f.inp_dt, f"Time the export was taken ({s['local_tz']}). "
         "Station times are compared after adding each station's offset (Stations tab)."),
        ("Threshold", "Late threshold (min)", s["threshold"], f.inp, "Late = actual - scheduled > this (Tableau RED)."),
        ("Grace", "Grace before 'No actual' (min)", s["grace"], f.inp,
         "Missing actual within this long after schedule = Pending (data entry / Tableau lag)."),
        ("EarlyCheck", "Early departure check (min)", s["early_check"], f.inp,
         "Departure more than this EARLY = Check date (keyed on the wrong day)."),
        ("LateCheck", "Late check (min)", s["late_check"], f.inp,
         "Departure or arrival more than this LATE = Check date."),
        ("MinTransit", "Minimum transit (min)", s["min_transit"], f.inp,
         "Arrival less than this after departure (or before it) = Check date."),
        ("MinLaneTrips", "Min clean trips for lane median", s["min_lane_trips"], f.inp,
         "Lanes tab: fewer clean trips than this = use the scheduled transit for Proj ATA."),
        ("WinFrom", "Report window from", s["win_from"], f.F(bg_color=INPUT, border=1, num_format="m/d/yyyy"),
         "Summary, Worklist and Report Card cover scheduled departure dates in this window."),
        ("WinTo", "Report window to", s["win_to"], f.F(bg_color=INPUT, border=1, num_format="m/d/yyyy"),
         "Data tab itself holds " + s["data_span"] + "."),
    ]
    ws.write_row(3, 0, ["Name", "Setting", "Value", "Meaning"], f.hv)
    for i, (name, lab, val, fmt, note) in enumerate(params, start=4):
        ws.write(i, 0, name, f.txt)
        ws.write(i, 1, lab, f.txt)
        if isinstance(val, (pd.Timestamp, datetime)):
            ws.write_datetime(i, 2, pd.Timestamp(val).to_pydatetime(), fmt)
        else:
            ws.write(i, 2, val, fmt)
        ws.write(i, 3, note, f.wrap)
        wb.define_name(name, f"=Settings!$C${i + 1}")
    r = 4 + len(params) + 2

    def table(r, title, headers, rows, name, note=None, widths=None):
        ws.write(r, 0, title, f.h2)
        if note:
            ws.write(r, 3, note, f.sub)
        ws.write_row(r + 1, 0, headers, f.hv)
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                ws.write(r + 2 + i, j, v, f.inp if j < len(headers) - 1 or name == "FaultTbl" else f.wrap)
        if name:
            wb.define_name(name, f"=Settings!$A${r + 3}:${col_letter(len(headers) - 1)}${r + 2 + len(rows)}")
        return r + len(rows) + 3

    r = table(r, "Truck types (by first 3 letters of AC Type)", ["Prefix", "Truck type", "Arrival rule", "Why"],
              [list(t) for t in TYPE_ROWS], "TypeTbl",
              "Arrival rule None = never measured; STA = keyed as STA, not measured. Others = Standard / Normal.")
    r = table(r, "Worklist issues", ["Issue", "Priority", "", "Meaning"],
              [(a, b, "", c) for a, b, c in ISSUES], "IssueTbl")
    r = table(r, "What to do (Worklist 'How to check')", ["Key", "", "", "Step ({DEST} = destination station)"],
              [(k, "", "", t) for k, t in STEPS], None)
    wb.define_name("StepsTbl", f"=Settings!$A${r - len(STEPS)}:$D${r - 1}")
    r = table(r, "Delay reasons -> Fault (Report Card)", ["Reason", "Fault"], [list(x) for x in REASONS], "FaultTbl")
    ws.write(r, 0, "Worklist status list", f.h2)
    for i, st in enumerate(STATUSES):
        ws.write(r + 1 + i, 0, st, f.inp)
    wb.define_name("StatusList", f"=Settings!$A${r + 2}:$A${r + 1 + len(STATUSES)}")
    wb.define_name("ReasonList", f"=Settings!$A${r - len(REASONS)}:$A${r - 1}")
    return ws


# ------------------------------------------------------------------------------------------------ stations
def write_stations(wb, f, df: pd.DataFrame, s: dict):
    ws = wb.add_worksheet("Stations")
    stations = sorted(set(df["origin"].dropna()) | set(df["dest"].dropna()) | set(OFFLINE_REGION))
    head(ws, f, "Stations - offline list and clock offsets", 10, s["as_of_txt"],
         "Offline = no iCargo access (Trucking Training Guide): get the ATA from FSU/RCF in OPR352 or email the "
         "station. Offset = hours the station clock is ahead of the as-of clock; used by Data[Dep/Arr Status].")
    hdr = ["Station", "Region", "Offline", "Note", "Time zone", "Offset (h)", "Out (window)", "In (window)",
           "Late dep (window)", "Late arr (window)", "Station email", "Station phone"]
    for j, h in enumerate(hdr):
        ws.write(3, j, h, f.hi if j >= 10 else f.hv if j < 6 else f.hf)
    contacts = s.get("contacts", {})
    tz = s["station_tz"]
    off = s["station_offset"]
    win = df[df["in_window"]]
    for i, st in enumerate(stations, start=4):
        r = i + 1
        ws.write(i, 0, st)
        ws.write(i, 1, OFFLINE_REGION.get(st, ""))
        ws.write(i, 2, "Y" if st in OFFLINE_REGION else "", f.inp)
        ws.write(i, 3, OFFLINE_NOTES.get(st, ""))
        ws.write(i, 4, tz.get(st, "") or "unknown - PC clock", f.txt)
        ws.write_number(i, 5, round(off.get(st, 0.0), 2), f.num1)
        ws.write_formula(i, 6, f'=COUNTIFS(Data[Origin],$A{r},Data[In Window],1)', f.num,
                         int((win["origin"] == st).sum()))
        ws.write_formula(i, 7, f'=COUNTIFS(Data[Dest],$A{r},Data[In Window],1)', f.num, int((win["dest"] == st).sum()))
        ws.write_formula(i, 8, f'=SUMIFS(Data[Late Dep],Data[Origin],$A{r},Data[In Window],1)', f.num,
                         int(((win["origin"] == st) & (win["dep_status"] == "Late")).sum()))
        ws.write_formula(i, 9, f'=SUMIFS(Data[Late Arr],Data[Dest],$A{r},Data[In Window],1)', f.num,
                         int(((win["dest"] == st) & (win["arr_status"] == "Late")).sum()))
        c = contacts.get(("station", st), {})
        ws.write(i, 10, c.get("Email", ""), f.inp)
        ws.write(i, 11, c.get("Phone", ""), f.inp)
    last = 4 + len(stations)
    wb.define_name("StationTbl", f"=Stations!$A$5:$L${last}")
    ws.set_column(0, 2, 8)
    ws.set_column(3, 3, 60)
    ws.set_column(4, 4, 20)
    ws.set_column(5, 9, 10)
    ws.set_column(10, 10, 30)
    ws.set_column(11, 11, 16)
    ws.freeze_panes(4, 1)
    ws.autofilter(3, 0, last - 1, len(hdr) - 1)


# ------------------------------------------------------------------------------------------------ vendors
def write_vendors(wb, f, df: pd.DataFrame, s: dict):
    ws = wb.add_worksheet("Vendors")
    head(ws, f, "Vendors - contacts and performance", 11, s["as_of_txt"],
         "Type each vendor's email once (yellow); the Worklist 'Email vendor' link uses it and it carries into the "
         "next report. Performance columns are SUMIFS on Data for the report window.")
    hdr = ["Vendor", "Vendor Code", "Email", "Phone", "Contact notes", "Trucks (window)", "Late dep", "Late dep %",
           "Late arr", "On-time arr %", "To work"]
    for j, h in enumerate(hdr):
        ws.write(3, j, h, f.hv if j < 2 else f.hi if j < 5 else f.hf)
    contacts = s.get("contacts", {})
    win = df[df["in_window"]]
    vend = (df.groupby("vendor")["vendor_code"].agg(lambda x: x.dropna().mode().iat[0] if x.notna().any() else "")
            .reset_index())
    order = win.groupby("vendor").size()
    vend["_n"] = vend["vendor"].map(order).fillna(0)
    vend = vend.sort_values(["_n", "vendor"], ascending=[False, True])
    for i, (v, code) in enumerate(zip(vend["vendor"], vend["vendor_code"]), start=4):
        r = i + 1
        g = win[win["vendor"] == v]
        c = contacts.get(("vendor", v), {})
        ws.write(i, 0, v)
        ws.write(i, 1, code or "")
        ws.write(i, 2, c.get("Email", ""), f.inp)
        ws.write(i, 3, c.get("Phone", ""), f.inp)
        ws.write(i, 4, c.get("Notes", ""), f.inp)
        crit = f"Data[Vendor],$A{r},Data[In Window],1"
        n = len(g)
        ld = int((g["dep_status"] == "Late").sum())
        la = int((g["arr_status"] == "Late").sum())
        am = int(g["arr_status"].isin(["Late", "On time"]).sum())
        ws.write_formula(i, 5, f"=COUNTIFS({crit})", f.num, n)
        ws.write_formula(i, 6, f"=SUMIFS(Data[Late Dep],{crit})", f.num, ld)
        ws.write_formula(i, 7, f'=IFERROR(G{r}/F{r},"")', f.pct, ld / n if n else "")
        ws.write_formula(i, 8, f"=SUMIFS(Data[Late Arr],{crit})", f.num, la)
        ws.write_formula(i, 9, f'=IFERROR(1-I{r}/SUMIFS(Data[Arr Measured],{crit}),"")', f.pct,
                         1 - la / am if am else "")
        ws.write_formula(i, 10, f"=SUMIFS(Data[To Work],{crit})", f.num, int((g["wl_issue"] != "").sum()))
    last = 4 + max(len(vend), 1)
    wb.define_name("VendorTbl", f"=Vendors!$A$5:$E${last}")
    ws.set_column(0, 0, 32)
    ws.set_column(1, 1, 12)
    ws.set_column(2, 2, 32)
    ws.set_column(3, 3, 16)
    ws.set_column(4, 4, 30)
    ws.set_column(5, 10, 10)
    ws.freeze_panes(4, 1)
    ws.autofilter(3, 0, last - 1, len(hdr) - 1)


def read_contacts(path: Path) -> dict:
    """Vendor / station emails typed into a previous workbook."""
    out = {}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return out
    try:
        for sheet, kind, key, fields in (("Vendors", "vendor", "Vendor", {"Email": "Email", "Phone": "Phone",
                                                                          "Contact notes": "Notes"}),
                                         ("Stations", "station", "Station", {"Station email": "Email",
                                                                             "Station phone": "Phone"})):
            if sheet not in wb.sheetnames:
                continue
            hdr = None
            for row in wb[sheet].iter_rows(values_only=True):
                if hdr is None:
                    if row and key in row:
                        hdr = {h: i for i, h in enumerate(row) if h}
                    continue
                k = row[hdr[key]] if hdr[key] < len(row) else None
                rec = {v: str(row[hdr[h]]) for h, v in fields.items()
                       if h in hdr and hdr[h] < len(row) and row[hdr[h]] not in (None, "")}
                if k and rec:
                    out[(kind, str(k))] = rec
    finally:
        wb.close()
    return out


def merge_contacts(store: Path, found: dict) -> dict:
    old = {}
    if store.exists():
        for rec in pd.read_csv(store, dtype=str, keep_default_na=False).to_dict("records"):
            old[(rec["Type"], rec["Key"])] = {k: rec[k] for k in ("Email", "Phone", "Notes") if rec.get(k)}
    for k, rec in found.items():
        old.setdefault(k, {}).update(rec)
    rows = [{"Type": t, "Key": key, "Email": r.get("Email", ""), "Phone": r.get("Phone", ""), "Notes": r.get("Notes", "")}
            for (t, key), r in sorted(old.items()) if any(r.values())]
    store.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows, columns=["Type", "Key", "Email", "Phone", "Notes"]).to_csv(store, index=False)
    return old


# ------------------------------------------------------------------------------------------------ lanes
def write_lanes(wb, f, df: pd.DataFrame, s: dict):
    ws = wb.add_worksheet("Lanes")
    head(ws, f, "Lanes - usual transit time", 6, s["as_of_txt"],
         "Median actual transit of clean trips (departure and arrival both Late/On time) in the Data tab. "
         "Used for Data[Proj ATA] when a lane has at least MinLaneTrips clean trips.")
    hdr = ["Lane", "Clean trips", "Median transit (min)", "Median transit (h)", "Trucks (window)", "Late (window)"]
    for j, h in enumerate(hdr):
        ws.write(3, j, h, f.hv if j == 0 else f.hf)
    ok = df["dep_status"].isin(["Late", "On time"]) & df["arr_status"].isin(["Late", "On time"])
    g = df[ok].groupby("lane")["transit_min"]
    cnt, med = g.size(), g.median()
    lanes = sorted(df["lane"].dropna().unique())
    for i, ln in enumerate(lanes, start=4):
        r = i + 1
        c = int(cnt.get(ln, 0))
        ws.write(i, 0, ln)
        ws.write_formula(i, 1, f'=COUNTIFS(Data[Lane],$A{r},Data[Arr Measured],1,Data[Dep Status],"Late")'
                               f'+COUNTIFS(Data[Lane],$A{r},Data[Arr Measured],1,Data[Dep Status],"On time")', f.num, c)
        m = float(med.get(ln)) if c >= s["min_lane_trips"] else ""
        ws.write_array_formula(i, 2, i, 2,
                               f'{{=IF($B{r}<MinLaneTrips,"",MEDIAN(IF(Data[Lane]=$A{r},IF(Data[Arr Measured]=1,'
                               f'IF((Data[Dep Status]="Late")+(Data[Dep Status]="On time"),Data[Transit (min)])))))}}',
                               f.num1, m)
        ws.write_formula(i, 3, f'=IF(C{r}="","",C{r}/60)', f.num1, (m / 60) if m != "" else "")
        w = df[df["in_window"] & (df["lane"] == ln)]
        ws.write_formula(i, 4, f'=COUNTIFS(Data[Lane],$A{r},Data[In Window],1)', f.num, len(w))
        ws.write_formula(i, 5, f'=SUMIFS(Data[Late Any],Data[Lane],$A{r},Data[In Window],1)', f.num,
                         int(w["late"].sum()))
    last = 4 + max(len(lanes), 1)
    wb.define_name("LaneTbl", f"=Lanes!$A$5:$D${last}")
    ws.set_column(0, 0, 24)
    ws.set_column(1, 5, 12)
    ws.freeze_panes(4, 1)


# ------------------------------------------------------------------------------------------------ summary
METRICS = [  # header, formula template ({c} = criteria, {X} = column letter of metric X in the same row), python
    ("Trucks", "=COUNTIFS({c})", lambda g: len(g), "int_"),
    ("Late dep", "=SUMIFS(Data[Late Dep],{c})", lambda g: int((g.dep_status == "Late").sum()), "int_"),
    ("Late dep %", '=IFERROR({Late dep}/{Trucks},"")', None, "pct"),
    ("Arr measured", "=SUMIFS(Data[Arr Measured],{c})", lambda g: int(g.arr_status.isin(["Late", "On time"]).sum()),
     "int_"),
    ("Late arr", "=SUMIFS(Data[Late Arr],{c})", lambda g: int((g.arr_status == "Late").sum()), "int_"),
    ("On-time arr %", '=IFERROR(1-{Late arr}/{Arr measured},"")', None, "pct"),
    ("Late (dep or arr)", "=SUMIFS(Data[Late Any],{c})", lambda g: int(g.late.sum()), "int_"),
    ("Avg min late (late arr)", '=IFERROR(AVERAGEIFS(Data[Arr Delay (min)],{c},Data[Late Arr],1),"")',
     lambda g: _mean(g.arr_delay[g.arr_status == "Late"]), "num"),
    ("No ATD", "=SUMIFS(Data[No ATD],{c})", lambda g: int((g.dep_status == "No actual").sum()), "int_"),
    ("No ATA", "=SUMIFS(Data[No ATA],{c})",
     lambda g: int(((g.arr_status == "No actual") & (g.dep_status != "No actual")).sum()), "int_"),
    ("Check date", "=SUMIFS(Data[Check Date],{c})",
     lambda g: int(((g.dep_status == "Check date") | (g.arr_status == "Check date")).sum()), "int_"),
]


def _mean(x):
    x = pd.to_numeric(x, errors="coerce").dropna()
    return float(x.mean()) if len(x) else ""


def _pyval(name, g, row_vals):
    for h, _, fn, _ in METRICS:
        if h == name:
            if fn:
                return fn(g)
    if name == "Late dep %":
        return row_vals["Late dep"] / row_vals["Trucks"] if row_vals["Trucks"] else ""
    if name == "On-time arr %":
        return 1 - row_vals["Late arr"] / row_vals["Arr measured"] if row_vals["Arr measured"] else ""
    return ""


def metric_row(ws, f, r, c0, crit, g, metrics=METRICS):
    """Write one row of metrics starting at column c0 (0-based); returns {metric: python value}."""
    pos = {h: col_letter(c0 + k) + str(r + 1) for k, (h, *_) in enumerate(metrics)}
    vals = {}
    for k, (h, tmpl, fn, fk) in enumerate(metrics):
        v = _pyval(h, g, vals)
        vals[h] = v
        formula = tmpl.replace("{c}", crit)
        for name, addr in pos.items():
            formula = formula.replace("{" + name + "}", addr)
        ws.write_formula(r, c0 + k, formula, getattr(f, fk), _v(v))
    return vals


def group_table(ws, f, r, title, keys, key_headers, crit_fn, sub_fn, note=None, extra_crit=",Data[In Window],1"):
    ws.write(r, 0, title, f.h2)
    if note:
        ws.write(r, len(key_headers) + 2, note, f.sub)
    for j, h in enumerate(key_headers):
        ws.write(r + 1, j, h, f.hv)
    for k, (h, *_) in enumerate(METRICS):
        ws.write(r + 1, len(key_headers) + k, h, f.hf)
    ws.set_row(r + 1, 30)
    first = r + 2
    for i, key in enumerate(keys):
        rr = first + i
        kk = key if isinstance(key, tuple) else (key,)
        for j, v in enumerate(kk):
            ws.write(rr, j, _v(v), f.date if isinstance(v, pd.Timestamp) else f.txt)
        metric_row(ws, f, rr, len(key_headers), crit_fn(rr + 1) + extra_crit, sub_fn(key))
    if not keys:
        ws.write(first, 0, "none", f.sub)
    return first, first + max(len(keys), 1) - 1


def write_summary(wb, f, df: pd.DataFrame, s: dict, pos: dict):
    ws = wb.add_worksheet("Summary")
    head(ws, f, "Summary - report window", 14, s["as_of_txt"],
         f"Scheduled departure {s['win_from']:%m/%d/%Y} to {s['win_to']:%m/%d/%Y} (Settings WinFrom/WinTo). Every "
         "cell is a COUNTIFS / SUMIFS on the Data tab with Data[In Window]=1 - click a cell to see it.")
    win = df[df["in_window"]]
    ws.write(3, 0, "Network", f.h2)
    for k, (h, *_) in enumerate(METRICS):
        ws.write(4, 1 + k, h, f.hf)
    ws.write(5, 0, "All trucks in window", f.txt)
    metric_row(ws, f, 5, 1, "Data[In Window],1", win)
    pos["kpi_row"] = 6
    pos["kpi_col"] = {h: col_letter(1 + k) for k, (h, *_) in enumerate(METRICS)}
    extra = []
    if s["has_rc"]:
        extra.append(("Late, not on Report Card", '=COUNTIFS(Data[In Window],1,Data[Late Any],1,Data[On Report Card],"<>Y")',
                      int((win["late"] & ~win["logged"]).sum())))
    extra.append(("Trucks to work (Worklist)", "=SUMIFS(Data[To Work],Data[In Window],1)",
                  int((win["wl_issue"] != "").sum())))
    for i, (lab, fo, v) in enumerate(extra):
        ws.write(7 + i, 0, lab, f.txt)
        ws.write_formula(7 + i, 1, fo, f.int_, v)
        pos["kpi_extra_" + str(i)] = f"Summary!$B${8 + i}"
    r = 7 + len(extra) + 2

    def by(field, label, title, top=None, sort="late"):
        nonlocal r
        g = win.groupby(field, dropna=True)
        order = (g["late"].sum().to_frame("l").join(g.size().to_frame("n")).sort_values(["l", "n"], ascending=False))
        keys = list(order.index[:top] if top else order.index)
        a, b = group_table(ws, f, r, title, keys, [label], lambda rr: f"Data[{field_hdr[field]}],$A{rr}",
                           lambda k: win[win[field] == k],
                           note=f"top {top} by late trucks" if top else "sorted by late trucks at build time")
        pos[f"by_{field}"] = (a, b)
        r = b + 3

    field_hdr = {"vendor": "Vendor", "origin": "Origin", "dest": "Dest", "lane": "Lane"}
    by("vendor", "Vendor", "By vendor")
    by("origin", "Origin", "By origin station")
    by("dest", "Dest", "By destination station")
    by("lane", "Lane", "By lane", top=40)
    # repeat late trucks
    lt = win[win["late"] & win["truck"].notna()]
    rep = lt.groupby(["truck", "lane"]).size()
    rep = rep[rep >= 3].sort_values(ascending=False)
    keys = list(rep.index)
    a, b = group_table(ws, f, r, "Repeat late trucks (late on 3+ days in the window)", keys, ["Truck", "Lane"],
                       lambda rr: f"Data[Truck],$A{rr},Data[Lane],$B{rr}",
                       lambda k: win[(win["truck"] == k[0]) & (win["lane"] == k[1])])
    r = b + 3
    days = sorted(win["sched_date"].dropna().unique(), reverse=True)
    a, b = group_table(ws, f, r, "By day", [pd.Timestamp(d) for d in days], ["Date"],
                       lambda rr: f"Data[Date],$A{rr}", lambda k: win[win["sched_date"] == k])
    ws.set_column(0, 0, 34)
    ws.set_column(1, 1, 12)
    ws.set_column(2, 14, 10)
    ws.freeze_panes(3, 1)
    return ws


# ------------------------------------------------------------------------------------------------ trends
def write_trends(wb, f, df: pd.DataFrame, s: dict, pos: dict):
    ws = wb.add_worksheet("Trends")
    head(ws, f, "Trends", 14, s["as_of_txt"],
         "Daily (last 35 full days) and weekly (all weeks in the Data tab) results, all trucks - COUNTIFS / SUMIFS on Data. "
         "The Dashboard charts read these tables.")
    end = min(pd.Timestamp(s["as_of"]).normalize(), df["sched_date"].max()) if len(df) else pd.Timestamp(s["as_of"])
    last_full = end - pd.Timedelta(days=1)       # the as-of day is still in progress
    days = [last_full - pd.Timedelta(days=i) for i in range(34, -1, -1)]
    a, b = group_table(ws, f, 3, "Daily (scheduled departure date, last 35 full days)", days, ["Date"],
                       lambda rr: f"Data[Date],$A{rr}", lambda k: df[df["sched_date"] == k], extra_crit="")
    pos["daily"] = (a, b)
    week = df["sched_date"] - pd.to_timedelta(df["sched_date"].dt.weekday, unit="D")
    weeks = sorted(w for w in week.dropna().unique() if pd.Timestamp(w) <= end)
    weeks = [pd.Timestamp(w) for w in weeks]
    r = b + 3
    a, b = group_table(ws, f, r, "Weekly (week starting Monday)", weeks, ["Week"],
                       lambda rr: f"Data[Week],$A{rr}", lambda k: df[week == k], extra_crit="",
                       note="the first and last week can be partial")
    pos["weekly"] = (a, b)
    r = b + 3
    # vendor x week late %
    top = df.groupby("vendor").size().sort_values(ascending=False).head(15).index
    ws.write(r, 0, "Late % by vendor and week (late dep or arr / trucks) - top 15 vendors by volume", f.h2)
    ws.write(r + 1, 0, "Vendor", f.hv)
    for k, w in enumerate(weeks):
        ws.write_datetime(r + 1, 1 + k, w.to_pydatetime(), f.F(bold=True, font_color="white", bg_color=FORM,
                                                               num_format="m/d"))
    for i, v in enumerate(top):
        rr = r + 2 + i
        ws.write(rr, 0, v)
        for k, w in enumerate(weeks):
            c = col_letter(1 + k)
            g = df[(df["vendor"] == v) & (week == w)]
            val = (g["late"].sum() / len(g)) if len(g) else ""
            ws.write_formula(rr, 1 + k, f'=IFERROR(SUMIFS(Data[Late Any],Data[Vendor],$A{rr + 1},Data[Week],{c}${r + 2})'
                                        f'/COUNTIFS(Data[Vendor],$A{rr + 1},Data[Week],{c}${r + 2}),"")', f.pct, _v(val))
    if len(top) and weeks:
        ws.conditional_format(r + 2, 1, r + 1 + len(top), len(weeks), {
            "type": "3_color_scale", "min_color": "#FFFFFF", "mid_color": "#FFE699", "max_color": RED})
    ws.set_column(0, 0, 30)
    ws.set_column(1, 30, 10)
    ws.freeze_panes(3, 1)
    return ws


# ------------------------------------------------------------------------------------------------ worklist
WL_COLS = [  # header, width, kind (f = formula, i = input, v = value)
    ("Priority", 6, "f"), ("Status", 15, "i"), ("Issue", 22, "f"), ("Truck", 8, "f"), ("Date", 9, "f"),
    ("iCargo (TRK001)", 15, "f"), ("Lane", 10, "f"), ("Vendor", 22, "f"), ("Truck Type", 8, "f"),
    ("Dest Offline", 8, "f"), ("STD", 12, "f"), ("ATD", 12, "f"), ("Dep Delay (min)", 7, "f"), ("STA", 12, "f"),
    ("ATA", 12, "f"), ("Arr Delay (min)", 7, "f"), ("Proj ATA", 12, "f"), ("How to check", 60, "f"),
    ("Email vendor", 12, "f"), ("Email station", 12, "f"),
    ("Vendor ETA", 13, "i"), ("Delay Reason", 10, "i"), ("Fault", 6, "f"), ("Updated in iCargo", 8, "i"),
    ("Notes", 40, "i"), ("Worked by", 10, "i"), ("Worked on", 10, "i"), ("On Report Card", 7, "f"),
    ("Inbox search", 14, "f"), ("Email subject", 34, "f"), ("Email body", 50, "f"),
    ("Step key", 13, "f"), ("TLM", 12, "v"), ("Data row", 6, "f"),
]
WL_INPUTS = ["Status", "Vendor ETA", "Delay Reason", "Updated in iCargo", "Notes", "Worked by", "Worked on"]


def worklist_rows(df: pd.DataFrame) -> pd.DataFrame:
    w = df[df["in_window"] & (df["wl_issue"] != "")].copy()
    pr = {a: b for a, b, _ in ISSUES}
    w["_p"] = w["wl_issue"].map(pr)
    return w.sort_values(["_p", "sched_date", "sched_dep"], ascending=[True, False, False])


def write_worklist(wb, f, df: pd.DataFrame, s: dict, notes: dict, pos: dict):
    ws = wb.add_worksheet("Worklist")
    w = worklist_rows(df)
    n = len(w)
    head(ws, f, f"Worklist - trucks to work ({s['win_from']:%m/%d} to {s['win_to']:%m/%d})", len(WL_COLS),
         s["as_of_txt"],
         "One row per truck with an open issue, worst first. Green = looked up from Data (live), yellow = yours. "
         "Work each in iCargo using 'How to check', then set Status. Your entries carry into tomorrow's report by TLM.")
    HR = 3
    c = {h: col_letter(j) for j, (h, *_) in enumerate(WL_COLS)}
    for j, (h, wdt, kind) in enumerate(WL_COLS):
        ws.write(HR, j, h, {"f": f.hf, "i": f.hi, "v": f.hv}[kind])
        ws.set_column(j, j, wdt)
    ws.set_row(HR, 32)
    steps = {k: t for k, t in STEPS}
    pr = {a: b for a, b, _ in ISSUES}
    fault = dict(REASONS)
    fm = {"Date": f.date, "STD": f.dtm, "ATD": f.dtm, "STA": f.dtm, "ATA": f.dtm, "Proj ATA": f.dtm,
          "Dep Delay (min)": f.num, "Arr Delay (min)": f.num, "Priority": f.num, "Data row": f.num,
          "How to check": f.F(text_wrap=True, valign="top", font_size=9),
          "Email body": f.F(text_wrap=True, valign="top", font_size=9), "Email subject": f.wrap}
    data_pos = {tlm: i + 1 for i, tlm in enumerate(df["tlm"])}
    contacts = s.get("contacts", {})
    link = f.F(font_color="#0563C1", underline=1)
    fm["Email vendor"] = link
    fm["Email station"] = link
    for i, t in enumerate(w.itertuples(index=False)):
        r = HR + 1 + i
        x = r + 1
        row = "$" + c["Data row"] + str(x)

        def D(col):
            return f"INDEX(Data[{col}],{row})"

        offline = "Offline" if t.dest in OFFLINE_REGION else ""
        key = ("NOATD" if t.wl_issue.startswith("No ATD") else
               ("NOATA_MAIL" if t.truck_type == "Mail" else "NOATA_OFFLINE" if offline else "NOATA")
               if t.wl_issue.startswith("No ATA") else "FIX" if t.wl_issue.startswith("Fix") else
               "PROJ" if t.wl_issue == "Projected late" else "LATE")
        note = notes.get(t.tlm, {})
        date_s = ddmmm(t.sched_date)
        body = (f"Hello, please send the ETA and the reason for the delay for AA Cargo truck {t.truck} {t.origin}-{t.dest} "
                f"scheduled to depart {_tx(t.sched_dep)} and arrive {_tx(t.sched_arr)}"
                + (f" (departed {_tx(t.act_dep)})" if pd.notna(t.act_dep) else "") + ". Thank you, AA Cargo Trucking")
        cells = {
            "Priority": (f'=IFERROR(VLOOKUP({c["Issue"]}{x},IssueTbl,2,FALSE),9)', pr.get(t.wl_issue, 9)),
            "Issue": (f"={D('Issue')}", t.wl_issue),
            "Truck": (f"={D('Truck')}", t.truck),
            "Date": (f"={D('Date')}", t.sched_date),
            "iCargo (TRK001)": (f'="TRK"&{c["Truck"]}{x}&"/"&UPPER(TEXT({c["Date"]}{x},"ddmmm"))', f"TRK{t.truck}/{date_s}"),
            "Lane": (f"={D('Lane')}", t.lane),
            "Vendor": (f'={D("Vendor")}&""', t.vendor or ""),
            "Truck Type": (f"={D('Truck Type')}", t.truck_type),
            "Dest Offline": (f'=IF(IFERROR(VLOOKUP({D("Dest")},StationTbl,3,FALSE),"")="Y","Offline","")', offline),
            "STD": (f'=IF({D("STD")}="","",{D("STD")})', t.sched_dep),
            "ATD": (f'=IF({D("ATD")}="","",{D("ATD")})', t.act_dep),
            "Dep Delay (min)": (f"={D('Dep Delay (min)')}", t.dep_delay),
            "STA": (f'=IF({D("STA")}="","",{D("STA")})', t.sched_arr),
            "ATA": (f'=IF({D("ATA")}="","",{D("ATA")})', t.act_arr),
            "Arr Delay (min)": (f"={D('Arr Delay (min)')}", t.arr_delay),
            "Proj ATA": (f"={D('Proj ATA')}", t.proj_arr),
            "Step key": (f'=IF({c["Issue"]}{x}="No ATD - possible no-show","NOATD",IF({c["Issue"]}{x}="No ATA - arrival overdue",'
                         f'IF({c["Truck Type"]}{x}="Mail","NOATA_MAIL",IF({c["Dest Offline"]}{x}="Offline","NOATA_OFFLINE","NOATA")),'
                         f'IF({c["Issue"]}{x}="Fix record - date keyed wrong","FIX",IF({c["Issue"]}{x}="Projected late","PROJ","LATE"))))',
                         key),
            "How to check": (f'=SUBSTITUTE(IFERROR(VLOOKUP({c["Step key"]}{x},StepsTbl,4,FALSE),""),"{{DEST}}",{D("Dest")})',
                             steps[key].replace("{DEST}", t.dest or "")),
            "Inbox search": (f'={c["Truck"]}{x}&" "&UPPER(TEXT({c["Date"]}{x},"ddmmm"))', f"{t.truck} {date_s}"),
            "Email vendor": (f'=IF(IFERROR(VLOOKUP({c["Vendor"]}{x},VendorTbl,3,FALSE),"")&""="","add email on Vendors",'
                             f'HYPERLINK("mailto:"&VLOOKUP({c["Vendor"]}{x},VendorTbl,3,FALSE)&"?subject="&'
                             f'_xlfn.ENCODEURL({c["Email subject"]}{x})&"&body="&_xlfn.ENCODEURL({c["Email body"]}{x}),'
                             f'"Email vendor"))',
                             "Email vendor" if contacts.get(("vendor", t.vendor), {}).get("Email") else "add email on Vendors"),
            "Email station": (f'=IF(IFERROR(VLOOKUP({D("Dest")},StationTbl,11,FALSE),"")&""="","",'
                              f'HYPERLINK("mailto:"&VLOOKUP({D("Dest")},StationTbl,11,FALSE)&"?subject="&'
                              f'_xlfn.ENCODEURL("AA Cargo truck "&{c["iCargo (TRK001)"]}{x}&" "&{c["Lane"]}{x}&" - actual arrival time")'
                              f'&"&body="&_xlfn.ENCODEURL("Hello, please send the actual arrival time (ATA) for AA Cargo truck "'
                              f'&{c["Truck"]}{x}&" "&{c["Lane"]}{x}&" scheduled to arrive "&UPPER(TEXT({c["STA"]}{x},"ddmmm hhmm"))'
                              f'&". Thank you, AA Cargo Trucking"),"Email "&{D("Dest")}))',
                              ("Email " + (t.dest or "")) if contacts.get(("station", t.dest), {}).get("Email") else ""),
            "Email subject": (f'="AA Cargo truck "&{c["iCargo (TRK001)"]}{x}&" "&{c["Lane"]}{x}&" - ETA and reason for delay"',
                              f"AA Cargo truck TRK{t.truck}/{date_s} {t.lane} - ETA and reason for delay"),
            "Email body": (f'="Hello, please send the ETA and the reason for the delay for AA Cargo truck "&{c["Truck"]}{x}'
                           f'&" "&{c["Lane"]}{x}&" scheduled to depart "&UPPER(TEXT({c["STD"]}{x},"ddmmm hhmm"))'
                           f'&" and arrive "&UPPER(TEXT({c["STA"]}{x},"ddmmm hhmm"))&IF({c["ATD"]}{x}="",""," (departed "'
                           f'&UPPER(TEXT({c["ATD"]}{x},"ddmmm hhmm"))&")")&". Thank you, AA Cargo Trucking"', body),
            "Fault": (f'=IF({c["Delay Reason"]}{x}="","",IFERROR(VLOOKUP({c["Delay Reason"]}{x},FaultTbl,2,FALSE),""))',
                      fault.get(note.get("Delay Reason", ""), "")),
            "On Report Card": (f'={D("On Report Card")}&""', "Y" if t.logged else ""),
            "Data row": (f'=MATCH({c["TLM"]}{x},Data[TLM],0)', data_pos.get(t.tlm, "")),
        }
        for j, (h, wdt, kind) in enumerate(WL_COLS):
            if kind == "f":
                fo, val = cells[h]
                ws.write_formula(r, j, fo, fm.get(h, f.txt), _v(val))
            elif kind == "v":
                ws.write_string(r, j, str(t.tlm))
            else:
                v = note.get(h, "")
                if h == "Status" and not v:
                    v = "Not started"
                if h in ("Vendor ETA", "Worked on") and v:
                    ts = pd.to_datetime(v, errors="coerce")
                    if pd.notna(ts):
                        ws.write_datetime(r, j, ts.to_pydatetime(), f.inp_dt)
                        continue
                ws.write(r, j, v, f.inp_wrap if h == "Notes" else f.inp)
    last = HR + max(n, 1)
    if n:
        ws.data_validation(HR + 1, _ci("Status"), last, _ci("Status"), {"validate": "list", "source": "=StatusList"})
        ws.data_validation(HR + 1, _ci("Delay Reason"), last, _ci("Delay Reason"),
                           {"validate": "list", "source": "=ReasonList"})
        ws.data_validation(HR + 1, _ci("Updated in iCargo"), last, _ci("Updated in iCargo"),
                           {"validate": "list", "source": ["Y", "N"]})
        rng = f"A{HR + 2}:{col_letter(len(WL_COLS) - 1)}{last + 1}"
        st = c["Status"]
        ws.conditional_format(rng, {"type": "formula", "criteria": f'=OR(${st}{HR + 2}="Updated in iCargo",'
                                    f'${st}{HR + 2}="Logged on Report Card",${st}{HR + 2}="Closed - no action")',
                                    "format": f.greyfill})
        ws.conditional_format(f"A{HR + 2}:A{last + 1}", {"type": "cell", "criteria": "==", "value": 1,
                                                          "format": f.F(bold=True, font_color="white", bg_color=RED)})
        ws.conditional_format(f"A{HR + 2}:A{last + 1}", {"type": "cell", "criteria": "==", "value": 2,
                                                          "format": f.F(bold=True, font_color="white", bg_color="#F28E2B")})
        ws.conditional_format(f"A{HR + 2}:A{last + 1}", {"type": "cell", "criteria": ">=", "value": 3,
                                                          "format": f.F(bold=True, bg_color="#F1CE63")})
    for r in range(HR + 1, last + 1):
        ws.set_row(r, 45)          # room for the How to check / email text (wrapped)
    ws.freeze_panes(HR + 1, 4)
    ws.autofilter(HR, 0, last, len(WL_COLS) - 1)
    pos["wl_rows"] = (HR + 2, last + 1)
    pos["wl_col"] = c
    return w


def _ci(h):
    return [x[0] for x in WL_COLS].index(h)


def _tx(ts):
    return "" if ts is None or pd.isna(ts) else pd.Timestamp(ts).strftime("%d%b %H%M").upper()


# ------------------------------------------------------------------------------------------------ report card
RC_COLUMNS = ["Vendor", "Vendor code ", "Month", "Departure Date", "TRK NBR", "ORG", "DEST",
              "Scheduled Departure", "Scheduled Arrival", "Actual Departure", "Actual Arrival",
              "Source", "Source", "GCX", "EPX", "EPXTC", "MAIL", "AWB #", "Reason", "Fault", "Summary",
              "Agent", "Agent completed by"]


def write_report_card(wb, f, df: pd.DataFrame, s: dict, notes: dict, pos: dict):
    ws = wb.add_worksheet("Report Card")
    statuses = card_status(df, s["threshold"])
    win = df[df["in_window"] & statuses.ne("")].copy()
    win["_card_status"] = statuses.loc[win.index]
    win["_card_order"] = win["_card_status"].map({v: i for i, v in enumerate(CARD_STATUSES)})
    win = win.sort_values(["_card_order", "logged", "sched_date", "truck"], ascending=[True, True, False, True])
    head(ws, f, "Late trucks - Report Card entries", len(RC_COLUMNS), s["as_of_txt"],
         "Linked to Data and to your Worklist entries (Reason, Notes -> Summary, Worked by -> Agent). "
         "Filter X to Ready to log and Already on Report Card to blank, then copy visible A:W; Paste Special > Values. "
         "Pending = watch only (amber); Check record = verify (red); grey = already logged. Rebuild to refresh rows/order.")
    HR = 3
    helper = ["TLM", "Dep Delay (min)", "Arr Delay (min)", "Late on", "Already on Report Card", "Data row",
              "Worklist row"]
    for j, h in enumerate(RC_COLUMNS):
        ws.write(HR, j, h, f.hf)
    ws.write(HR, 23, "Card status", f.hf)
    ws.set_column(23, 23, 29)
    for i, status in enumerate(CARD_STATUSES):
        counts = int((win["_card_status"] == status).sum())
        ws.write_formula(2, i * 4, f'="{status}: "&COUNTIF(X5:X{max(len(win) + 4, 5)},"{status}")',
                         f.sub, f"{status}: {counts}")
    hx = len(RC_COLUMNS) + 1
    for k, h in enumerate(helper):
        ws.write(HR, hx + k, h, f.hv if h == "TLM" else f.hf)
    ws.set_row(HR, 30)
    widths = [26, 9, 6, 10, 8, 6, 6, 10, 10, 10, 10, 9, 8, 5, 5, 6, 5, 14, 9, 6, 40, 9, 9]
    for j, wdt in enumerate(widths):
        ws.set_column(j, j, wdt)
    ws.set_column(hx, hx + len(helper) - 1, 10)
    ws.write(HR, 31, "Proj ATA", f.hf)
    ws.write(HR, 32, "Proj Arr Delay (min)", f.hf)
    ws.set_column(31, 31, 17)
    ws.set_column(32, 32, 13)
    wl = pos["wl_col"]
    wl_first, wl_last = pos["wl_rows"]
    data_pos = {tlm: i + 1 for i, tlm in enumerate(df["tlm"])}
    wl_tlms = list(worklist_rows(df)["tlm"])
    fault = dict(REASONS)
    dcol, wcol = col_letter(hx + 5), col_letter(hx + 6)
    for i, t in enumerate(win.itertuples(index=False)):
        r = HR + 1 + i
        x = r + 1

        def D(col):
            return f"INDEX(Data[{col}],${dcol}{x})"

        def W(col):
            return (f'IF(${wcol}{x}="","",INDEX(Worklist!${wl[col]}${wl_first}:${wl[col]}${wl_last},${wcol}{x})&"")')

        def rct(col):
            return (f'IF({D(col)}="","",TEXT({D(col)},"hhmm")&IF(INT({D(col)})-{D("Date")}>0,"+"&(INT({D(col)})-{D("Date")}),'
                    f'IF(INT({D(col)})-{D("Date")}<0,"-"&({D("Date")}-INT({D(col)})),"")))')

        note = notes.get(t.tlm, {})
        base = t.sched_date
        dep_txt = "NO SHOW?" if t.dep_status == "No actual" else rc_time(t.act_dep, base)
        cells = [
            (f'={D("Vendor")}&""', t.vendor or "", f.txt),
            (f'={D("Vendor Code")}&""', t.vendor_code or "", f.txt),
            (f'=TEXT({D("Date")},"mmm")', f"{base:%b}", f.txt),
            (f"={D('Date')}", t.sched_date, f.date),
            (f'={D("Truck")}&""', t.truck or "", f.txt),
            (f'={D("Origin")}&""', t.origin or "", f.txt),
            (f'={D("Dest")}&""', t.dest or "", f.txt),
            ("=" + rct("STD"), rc_time(t.sched_dep, base), f.txt),
            ("=" + rct("STA"), rc_time(t.sched_arr, base), f.txt),
            (f'=IF({D("Dep Status")}="No actual","NO SHOW?",{rct("ATD")})', dep_txt, f.txt),
            ("=" + rct("ATA"), rc_time(t.act_arr, base), f.txt),
            (None, "TABLEAU", f.txt), (None, "", f.txt), (None, "", f.inp), (None, "", f.inp), (None, "", f.inp),
            (None, "", f.inp), (None, "", f.inp),
            ("=" + W("Delay Reason"), note.get("Delay Reason", ""), f.txt),
            (f'=IF(S{x}="","",IFERROR(VLOOKUP(S{x},FaultTbl,2,FALSE),""))', fault.get(note.get("Delay Reason", ""), ""),
             f.txt),
            ("=" + W("Notes"), note.get("Notes", ""), f.txt),
            ("=" + W("Worked by"), note.get("Worked by", ""), f.txt),
            (None, "", f.inp),
        ]
        for j, (fo, val, fmt) in enumerate(cells):
            if fo:
                ws.write_formula(r, j, fo, fmt, _v(val))
            else:
                ws.write(r, j, val, fmt)
        ws.write_formula(r, 23, f'={D("Card status")}&""', f.txt, statuses.loc[win.index[i]])
        ws.write_formula(r, 31, f'=IF({D("Proj ATA")}="","",{D("Proj ATA")})', f.dtm, _v(t.proj_arr))
        ws.write_formula(r, 32, f'=IF({D("Proj Arr Delay (min)")}="","",{D("Proj Arr Delay (min)")})',
                         f.num, _v(t.proj_arr_delay))
        on = ", ".join(a for a, st in (("Departure", t.dep_status), ("Arrival", t.arr_status)) if st == "Late")
        ws.write_string(r, hx, str(t.tlm))
        ws.write_formula(r, hx + 1, f"={D('Dep Delay (min)')}", f.num, _v(t.dep_delay))
        ws.write_formula(r, hx + 2, f"={D('Arr Delay (min)')}", f.num, _v(t.arr_delay))
        ws.write_formula(r, hx + 3, f'=IF(AND({D("Late Dep")}=1,{D("Late Arr")}=1),"Departure, Arrival",'
                                    f'IF({D("Late Dep")}=1,"Departure",IF({D("Late Arr")}=1,"Arrival","Not late yet")))',
                         f.txt, on or "Not late yet")
        ws.write_formula(r, hx + 4, f'={D("On Report Card")}&""', f.txt, "Y" if t.logged else "")
        ws.write_formula(r, hx + 5, f'=MATCH({col_letter(hx)}{x},Data[TLM],0)', f.num, data_pos.get(t.tlm, ""))
        wr = (wl_tlms.index(t.tlm) + 1) if t.tlm in wl_tlms else ""
        ws.write_formula(r, hx + 6, f'=IFERROR(MATCH({col_letter(hx)}{x},Worklist!${wl["TLM"]}${wl_first}:'
                                    f'${wl["TLM"]}${wl_last},0),"")', f.num, wr)
    last = HR + max(len(win), 1)
    if len(win):
        for criteria, fmt in ((f'=$X{HR + 2}="Check record"', f.red),
                              (f'=LEFT($X{HR + 2},7)="Pending"', f.amber)):
            ws.conditional_format(f"A{HR + 2}:AG{last + 1}", {
                "type": "formula", "criteria": criteria, "format": fmt, "stop_if_true": True})
        ws.conditional_format(f"A{HR + 2}:W{last + 1}", {"type": "formula",
                                                          "criteria": f'=${col_letter(hx + 4)}{HR + 2}="Y"',
                                                          "format": f.greyfill})
        lo = col_letter(hx + 3)
        ws.conditional_format(f"J{HR + 2}:J{last + 1}", {"type": "formula", "format": f.red,
                                                          "criteria": f'=OR(J{HR + 2}="NO SHOW?",'
                                                                      f'ISNUMBER(SEARCH("Departure",${lo}{HR + 2})))'})
        ws.conditional_format(f"K{HR + 2}:K{last + 1}", {"type": "formula", "format": f.red,
                                                          "criteria": f'=ISNUMBER(SEARCH("Arrival",${lo}{HR + 2}))'})
    ws.freeze_panes(HR + 1, 5)
    ws.autofilter(HR, 0, last, 32)


# ------------------------------------------------------------------------------------------------ dashboard
def write_dashboard(wb, ws, f, df: pd.DataFrame, s: dict, pos: dict, wl: pd.DataFrame, notes: dict):
    ws.hide_gridlines(2)
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    ws.set_column(0, 0, 2)
    ws.set_column(1, 16, 11.5)
    ws.merge_range(0, 1, 0, 12, "Late Trucks Dashboard", f.title)
    ws.merge_range(0, 13, 0, 16, f"As of {s['as_of_txt']}", f.title_r)
    ws.set_row(0, 30)
    ws.write(1, 1, f"Report window {s['win_from']:%m/%d/%Y} - {s['win_to']:%m/%d/%Y}. Tiles and charts are formulas "
                   "on Summary / Trends / Worklist, which are formulas on Data.", f.sub)
    kr = pos["kpi_row"]
    kc = pos["kpi_col"]
    win = df[df["in_window"]]
    meas = win["arr_status"].isin(["Late", "On time"]).sum()
    tiles = [
        ("Trucks", f"=Summary!{kc['Trucks']}{kr}", len(win), f.tile_v),
        ("Late departure %", f"=Summary!{kc['Late dep %']}{kr}", (win["dep_status"] == "Late").mean() if len(win) else "",
         f.tile_p),
        ("On-time arrival %", f"=Summary!{kc['On-time arr %']}{kr}",
         1 - (win["arr_status"] == "Late").sum() / meas if meas else "", f.tile_p),
        ("Late trucks", f"=Summary!{kc['Late (dep or arr)']}{kr}", int(win["late"].sum()), f.tile_r),
        ("No ATD", f"=Summary!{kc['No ATD']}{kr}", int((win["dep_status"] == "No actual").sum()), f.tile_r),
        ("No ATA", f"=Summary!{kc['No ATA']}{kr}",
         int(((win["arr_status"] == "No actual") & (win["dep_status"] != "No actual")).sum()), f.tile_r),
        ("Check date", f"=Summary!{kc['Check date']}{kr}",
         int(((win["dep_status"] == "Check date") | (win["arr_status"] == "Check date")).sum()), f.tile_v),
    ]
    for k, (lab, fo, val, fmt) in enumerate(tiles):
        c = 1 + k * 2
        ws.merge_range(3, c, 3, c + 1, lab, f.tile_h)
        ws.merge_range(4, c, 5, c + 1, "", fmt)
        ws.write_formula(4, c, fo, fmt, _v(val))
    ws.set_row(3, 20)
    ws.set_row(4, 22)
    ws.set_row(5, 22)

    # worklist progress tables (right side helper area, also visible)
    wc = pos["wl_col"]
    a, b = pos["wl_rows"]
    st_rng = f"Worklist!${wc['Status']}${a}:${wc['Status']}${b}"
    is_rng = f"Worklist!${wc['Issue']}${a}:${wc['Issue']}${b}"
    pr_rng = f"Worklist!${wc['Priority']}${a}:${wc['Priority']}${b}"
    st_now = pd.Series([notes.get(t, {}).get("Status") or "Not started" for t in wl["tlm"]], dtype=object)
    R0 = 7
    ws.write(R0, 1, "Worklist by issue", f.h2)
    ws.merge_range(R0 + 1, 1, R0 + 1, 3, "Issue", f.hf)
    ws.write_row(R0 + 1, 4, ["Trucks", "Not started", "Done"], f.hf)
    for i, (iss, p, _) in enumerate(ISSUES):
        r = R0 + 2 + i
        m = (wl["wl_issue"] == iss).values
        ws.merge_range(r, 1, r, 3, iss, f.txt)
        ws.write_formula(r, 4, f'=COUNTIF({is_rng},B{r + 1})', f.num, int(m.sum()))
        ws.write_formula(r, 5, f'=COUNTIFS({is_rng},B{r + 1},{st_rng},"Not started")', f.num,
                         int((m & (st_now == "Not started").values).sum()))
        done = "+".join(f'COUNTIFS({is_rng},B{r + 1},{st_rng},"{d}")' for d in DONE)
        ws.write_formula(r, 6, "=" + done, f.num, int((m & st_now.isin(DONE).values).sum()))
    r_tot = R0 + 2 + len(ISSUES)
    ws.write(r_tot, 1, "Total", f.F(bold=True))
    tot_v = {4: len(wl), 5: int((st_now == "Not started").sum()), 6: int(st_now.isin(DONE).sum())}
    for j in (4, 5, 6):
        cl = col_letter(j)
        ws.write_formula(r_tot, j, f"=SUM({cl}{R0 + 3}:{cl}{r_tot})", f.F(bold=True, num_format="0"), tot_v[j])
    ws.write(R0, 9, "Worklist by status", f.h2)
    ws.merge_range(R0 + 1, 9, R0 + 1, 10, "Status", f.hf)
    ws.write(R0 + 1, 11, "Trucks", f.hf)
    for i, stt in enumerate(STATUSES):
        r = R0 + 2 + i
        ws.merge_range(r, 9, r, 10, stt, f.txt)
        ws.write_formula(r, 11, f'=COUNTIF({st_rng},J{r + 1})', f.num, int((st_now == stt).sum()))
    ws.write(R0, 13, "Priority 1 not started", f.h2)
    ws.write_formula(R0 + 1, 13, f'=COUNTIFS({pr_rng},1,{st_rng},"Not started")', f.tile_r,
                     int(((wl["wl_issue"].str.startswith("No ")).values & (st_now == "Not started").values).sum()))

    # charts
    def chart_ts(title, sheet, rows, cat_col, series, y_pct=False, combo=None):
        a, b = rows
        ch = wb.add_chart({"type": "column" if combo else "line"})
        if combo:
            name, col, color = combo
            ch.add_series({"name": name, "categories": [sheet, a - 1, cat_col, b - 1, cat_col],
                           "values": [sheet, a - 1, col, b - 1, col], "fill": {"color": color}, "gap": 60})
            ln = wb.add_chart({"type": "line"})
            for nm, cc, color in series:
                ln.add_series({"name": nm, "categories": [sheet, a - 1, cat_col, b - 1, cat_col],
                               "values": [sheet, a - 1, cc, b - 1, cc], "y2_axis": True,
                               "line": {"color": color, "width": 2.25}, "marker": {"type": "none"}})
            ch.combine(ln)
            ln.set_y2_axis({"num_format": "0%", "min": 0, "max": 1, "major_gridlines": {"visible": False}})
            ch.set_y_axis({"name": name, "major_gridlines": {"visible": True, "line": {"color": "#E7E6E6"}}})
        else:
            for nm, cc, color in series:
                ch.add_series({"name": nm, "categories": [sheet, a - 1, cat_col, b - 1, cat_col],
                               "values": [sheet, a - 1, cc, b - 1, cc], "line": {"color": color, "width": 2.25},
                               "marker": {"type": "circle", "size": 4, "fill": {"color": color},
                                          "border": {"color": color}}})
            ch.set_y_axis({"num_format": "0%" if y_pct else "0", "min": 0,
                           "major_gridlines": {"visible": True, "line": {"color": "#E7E6E6"}}})
        ch.set_title({"name": title, "name_font": {"size": 11, "bold": True}})
        ch.set_legend({"position": "bottom"})
        ch.set_x_axis({"num_format": "m/d", "num_font": {"rotation": -45, "size": 8}})
        ch.set_chartarea({"border": {"none": True}})
        return ch

    mi = {h: 1 + k for k, (h, *_) in enumerate(METRICS)}
    d_rows = (pos["daily"][0] + 1, pos["daily"][1] + 1)
    w_rows = (pos["weekly"][0] + 1, pos["weekly"][1] + 1)
    c1 = chart_ts("Daily: trucks, late departure % and on-time arrival % (last 35 full days)", "Trends", d_rows, 0,
                  [("Late dep %", mi["Late dep %"], RED), ("On-time arr %", mi["On-time arr %"], GREEN)],
                  combo=("Trucks", mi["Trucks"], "#BDD7EE"))
    ws.insert_chart(20, 1, c1, {"x_scale": 1.55, "y_scale": 1.15})
    c2 = chart_ts("Weekly: late departure % and on-time arrival %", "Trends", w_rows, 0,
                  [("Late dep %", mi["Late dep %"], RED), ("On-time arr %", mi["On-time arr %"], GREEN)], y_pct=True)
    ws.insert_chart(20, 9, c2, {"x_scale": 1.25, "y_scale": 1.15})
    c3 = chart_ts("Daily: trucks needing follow-up (No ATD, No ATA, Check date)", "Trends", d_rows, 0,
                  [("No ATD", mi["No ATD"], "#7F7F7F"), ("No ATA", mi["No ATA"], "#F28E2B"),
                   ("Check date", mi["Check date"], AMBER)])
    ws.insert_chart(38, 1, c3, {"x_scale": 1.55, "y_scale": 1.0})

    def bar(title, sheet, rows, key_col, val_col, color, n=10):
        a, b = rows
        b = min(b, a + n - 1)
        ch = wb.add_chart({"type": "bar"})
        ch.add_series({"name": title, "categories": [sheet, a, key_col, b, key_col],
                       "values": [sheet, a, val_col, b, val_col], "fill": {"color": color}, "gap": 40,
                       "data_labels": {"value": True, "font": {"size": 8}}})
        ch.set_title({"name": title, "name_font": {"size": 11, "bold": True}})
        ch.set_legend({"none": True})
        ch.set_y_axis({"reverse": True, "num_font": {"size": 8}})
        ch.set_x_axis({"major_gridlines": {"visible": False}})
        ch.set_chartarea({"border": {"none": True}})
        return ch

    if "by_vendor" in pos:
        ws.insert_chart(38, 9, bar("Vendors with the most late trucks (window)", "Summary", pos["by_vendor"], 0,
                                   mi["Late (dep or arr)"], RED), {"x_scale": 1.25, "y_scale": 1.0})
    if "by_origin" in pos:
        ws.insert_chart(54, 1, bar("Origins with the most late trucks: late departures", "Summary",
                                   pos["by_origin"], 0, mi["Late dep"], "#F28E2B"), {"x_scale": 1.0, "y_scale": 1.0})
    if "by_dest" in pos:
        ws.insert_chart(54, 6, bar("Destinations with the most late trucks: late arrivals", "Summary",
                                   pos["by_dest"], 0, mi["Late arr"], "#4E79A7"), {"x_scale": 1.0, "y_scale": 1.0})
    ch = wb.add_chart({"type": "bar", "subtype": "stacked"})
    ch.add_series({"name": "Not started", "categories": ["Dashboard", R0 + 2, 1, R0 + 1 + len(ISSUES), 1],
                   "values": ["Dashboard", R0 + 2, 5, R0 + 1 + len(ISSUES), 5], "fill": {"color": RED}})
    ch.add_series({"name": "Done", "categories": ["Dashboard", R0 + 2, 1, R0 + 1 + len(ISSUES), 1],
                   "values": ["Dashboard", R0 + 2, 6, R0 + 1 + len(ISSUES), 6], "fill": {"color": GREEN}})
    ch.set_title({"name": "Worklist progress by issue", "name_font": {"size": 11, "bold": True}})
    ch.set_y_axis({"reverse": True, "num_font": {"size": 8}})
    ch.set_legend({"position": "bottom"})
    ch.set_chartarea({"border": {"none": True}})
    ws.insert_chart(54, 11, ch, {"x_scale": 1.0, "y_scale": 1.0})
    ws.activate()


# ------------------------------------------------------------------------------------------------ about
def write_about(wb, f, cols, s: dict):
    ws = wb.add_worksheet("About")
    ws.set_column(0, 0, 24)
    ws.set_column(1, 1, 10)
    ws.set_column(2, 2, 120)
    head(ws, f, "How the numbers are built (audit trail)", 3, s["as_of_txt"])
    lines = [
        ("1. Export", "", "Tableau custom-SQL export (" + s["sources"] + "), cleaned by late_trucks.py: one row per TLM "
                          "(superseded I<>M versions and duplicate legs dropped, TLMs with no flight leg excluded like "
                          "Tableau's Exclude Nulls). Times kept to the minute."),
        ("2. Data tab", "", f"{s['n_data']:,} trucks ({s['data_span']}). Blue columns are the export values; green columns "
                            "are formulas, listed below. Filter Data to see the exact trucks behind any number."),
        ("3. Settings", "", "Named cells AsOf, Threshold, Grace, EarlyCheck, LateCheck, MinTransit, MinLaneTrips, "
                            "WinFrom, WinTo and the lookup tables. Change one and the workbook re-scores."),
        ("4. Summary / Trends", "", "COUNTIFS / SUMIFS / AVERAGEIFS over Data columns (window rows use Data[In Window]=1)."),
        ("5. Lanes / Vendors / Stations", "", "Lane median transit (array MEDIAN over Data) feeds Data[Proj ATA]; station "
                                              "offline flag and clock offset feed Data statuses and the Worklist; vendor and "
                                              "station emails (yours) feed the Worklist email links."),
        ("6. Worklist", "", "Rows = trucks with Data[To Work]=1 at build time; every green cell is INDEX(Data[...], Data "
                            "row). Yellow cells are yours and carry forward to the next report by TLM."),
        ("7. Report Card", "", "INDEX into Data plus your Worklist Reason / Notes / Worked by."),
        ("8. Dashboard", "", "Tiles point at Summary; charts read Trends / Summary; progress tables COUNTIF the Worklist."),
        ("Check", "", s.get("check_txt", "")),
        ("", "", ""),
        ("Data column", "Type", "Formula / source"),
    ]
    for i, (a, b, c) in enumerate(lines, start=3):
        bold = a in ("Data column",)
        ws.write(i, 0, a, f.hv if bold else f.F(bold=True, valign="top"))
        ws.write(i, 1, b, f.hv if bold else f.txt)
        ws.write(i, 2, c, f.hv if bold else f.wrap)
    r = 3 + len(lines)
    for h, kind, _, _, formula, _ in cols:
        ws.write(r, 0, h, f.txt)
        ws.write(r, 1, "formula" if kind == "f" else "export", f.txt)
        ws.write_string(r, 2, (formula or "value from the export").replace("Data[[#This Row],[", "[@["), f.wrap)
        r += 1


# ------------------------------------------------------------------------------------------------ daily book
def build_daily_book(df: pd.DataFrame, path: Path, s: dict, notes: dict) -> dict:
    """df = Data rows (already scored by late_trucks.py, with in_window, wl_issue, logged, now_o/now_d)."""
    import xlsxwriter
    _speed_up_xlsxwriter()
    df = df.sort_values(["sched_date", "sched_dep", "truck"], ascending=[False, False, True]).reset_index(drop=True)
    for c in ("sched_dep", "sched_arr", "act_dep", "act_arr"):
        df[c] = df[c].dt.floor("min")
    wb = xlsxwriter.Workbook(str(path), {"nan_inf_to_errors": True})
    wb.set_calc_mode("auto")
    f = Fmt(wb)
    pos = {}
    # sheet order = creation order; the Dashboard is created first and filled last (it needs the others' positions)
    dash = wb.add_worksheet("Dashboard")
    wl_df = write_worklist(wb, f, df, s, notes, pos)
    write_report_card(wb, f, df, s, notes, pos)
    write_summary(wb, f, df, s, pos)
    write_trends(wb, f, df, s, pos)
    cols = write_data(wb, f, df, s)
    write_lanes(wb, f, df, s)
    write_vendors(wb, f, df, s)
    write_stations(wb, f, df, s)
    write_settings(wb, f, s)
    write_about(wb, f, cols, s)
    write_dashboard(wb, dash, f, df, s, pos, wl_df, notes)
    if s.get("shift_comparison") is not None:
        from shift_compare import write_variance
        write_variance(wb, s["shift_comparison"])
    wb.close()
    return {"worklist": wl_df, "pos": pos}


# ------------------------------------------------------------------------------------------------ notes
def read_notes(path: Path) -> dict:
    """Coordinator inputs from a previous workbook's Worklist (or a legacy Trucks Report sheet), keyed by TLM."""
    out = {}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return out
    try:
        if "Worklist" in wb.sheetnames:
            ws = wb["Worklist"]
            hdr = None
            for row in ws.iter_rows(values_only=True):
                if hdr is None:
                    if row and "TLM" in row and "Status" in row:
                        hdr = {h: i for i, h in enumerate(row) if h}
                    continue
                tlm = row[hdr["TLM"]] if hdr["TLM"] < len(row) else None
                if not tlm:
                    continue
                rec = {}
                for h in WL_INPUTS:
                    if h in hdr and hdr[h] < len(row) and row[hdr[h]] not in (None, ""):
                        v = row[hdr[h]]
                        rec[h] = v.strftime("%Y-%m-%d %H:%M") if hasattr(v, "strftime") else str(v)
                if rec.get("Status") == "Not started":
                    rec.pop("Status")
                if rec:
                    out[str(tlm)] = rec
        elif "Trucks Report" in wb.sheetnames:  # v3 report: free text typed to the right of the table
            ws = wb["Trucks Report"]
            hdr = None
            for row in ws.iter_rows(values_only=True):
                if hdr is None:
                    if row and "TLM" in row:
                        hdr = row.index("TLM")
                        width = max(i for i, h in enumerate(row) if h) + 1
                    continue
                if hdr >= len(row) or not row[hdr]:
                    continue
                extra = [str(v) for v in row[width:] if v not in (None, "")]
                if extra:
                    out[str(row[hdr])] = {"Notes": "; ".join(extra)}
    finally:
        wb.close()
    return out


def merge_notes(store: Path, found: dict, as_of: datetime) -> dict:
    """Keep a running notes file (worklist_notes.csv) so entries survive between reports."""
    cols = ["TLM"] + WL_INPUTS + ["Updated"]
    old = {}
    if store.exists():
        d = pd.read_csv(store, dtype=str, keep_default_na=False)
        for rec in d.to_dict("records"):
            old[rec["TLM"]] = {k: v for k, v in rec.items() if k in WL_INPUTS and v}
    for tlm, rec in found.items():
        cur = old.setdefault(tlm, {})
        for k, v in rec.items():
            if v not in (None, ""):
                cur[k] = v
        cur["_upd"] = as_of.strftime("%Y-%m-%d %H:%M")
    rows = [{"TLM": t, **{k: r.get(k, "") for k in WL_INPUTS}, "Updated": r.get("_upd", "")} for t, r in old.items()
            if any(r.get(k) for k in WL_INPUTS)]
    store.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows, columns=cols).to_csv(store, index=False)
    return {t: {k: v for k, v in r.items() if k in WL_INPUTS} for t, r in old.items()}


# ------------------------------------------------------------------------------------------------ history book
HIST_COLS = [("TLM", "tlm", 12, "txt"), ("Date", "sched_date", 10, "date"), ("Month", "month", 8, "txt"),
             ("Week", "week", 10, "date"), ("Truck", "truck", 9, "txt"), ("Vendor", "vendor", 26, "txt"),
             ("Origin", "origin", 7, "txt"), ("Dest", "dest", 7, "txt"), ("Lane", "lane", 10, "txt"),
             ("AC Type", "ac_type", 7, "txt"), ("Truck Type", "truck_type", 9, "txt"),
             ("STD", "sched_dep", 13, "dtm"), ("ATD", "act_dep", 13, "dtm"), ("Dep Delay (min)", "dep_delay", 8, "num"),
             ("Dep Status", "dep_status", 10, "txt"), ("STA", "sched_arr", 13, "dtm"), ("ATA", "act_arr", 13, "dtm"),
             ("Arr Delay (min)", "arr_delay", 8, "num"), ("Arr Status", "arr_status", 10, "txt"),
             ("Late Dep", "f_ld", 5, "num"), ("Late Arr", "f_la", 5, "num"), ("Arr Measured", "f_am", 6, "num"),
             ("Late Any", "f_l", 5, "num"), ("Check Date", "f_cd", 5, "num"), ("No ATD", "f_nd", 5, "num"),
             ("No ATA", "f_na", 5, "num")]
H_METRICS = [
    ("Trucks", "=COUNTIFS({c})", lambda g: len(g), "int_"),
    ("Late dep", "=SUMIFS(Data[Late Dep],{c})", lambda g: int(g.f_ld.sum()), "int_"),
    ("Late dep %", '=IFERROR({Late dep}/{Trucks},"")', lambda g: g.f_ld.sum() / len(g) if len(g) else "", "pct"),
    ("Arr measured", "=SUMIFS(Data[Arr Measured],{c})", lambda g: int(g.f_am.sum()), "int_"),
    ("Late arr", "=SUMIFS(Data[Late Arr],{c})", lambda g: int(g.f_la.sum()), "int_"),
    ("On-time arr %", '=IFERROR(1-{Late arr}/{Arr measured},"")',
     lambda g: 1 - g.f_la.sum() / g.f_am.sum() if g.f_am.sum() else "", "pct"),
    ("Late (dep or arr)", "=SUMIFS(Data[Late Any],{c})", lambda g: int(g.f_l.sum()), "int_"),
    ("Late %", '=IFERROR({Late (dep or arr)}/{Trucks},"")', lambda g: g.f_l.sum() / len(g) if len(g) else "", "pct"),
    ("No ATD", "=SUMIFS(Data[No ATD],{c})", lambda g: int(g.f_nd.sum()), "int_"),
    ("No ATA", "=SUMIFS(Data[No ATA],{c})", lambda g: int(g.f_na.sum()), "int_"),
    ("Check date", "=SUMIFS(Data[Check Date],{c})", lambda g: int(g.f_cd.sum()), "int_"),
]


def build_history_book(df: pd.DataFrame, path: Path, as_of: datetime, months: int = 24) -> Path:
    import xlsxwriter
    _speed_up_xlsxwriter()
    end = pd.Timestamp(as_of).normalize()
    start = (end - pd.DateOffset(months=months - 1)).replace(day=1)
    x = df[(df["sched_date"] >= start) & (df["sched_date"] <= end)].copy()
    x = x.sort_values(["sched_date", "truck"]).reset_index(drop=True)
    for c in ("sched_dep", "sched_arr", "act_dep", "act_arr"):
        x[c] = x[c].dt.floor("min")
    x["month"] = x["sched_date"].dt.strftime("%Y-%m")
    x["week"] = x["sched_date"] - pd.to_timedelta(x["sched_date"].dt.weekday, unit="D")
    x["f_ld"] = (x["dep_status"] == "Late").astype(int)
    x["f_la"] = (x["arr_status"] == "Late").astype(int)
    x["f_am"] = x["arr_status"].isin(["Late", "On time"]).astype(int)
    x["f_l"] = x["late"].astype(int)
    x["f_cd"] = ((x["dep_status"] == "Check date") | (x["arr_status"] == "Check date")).astype(int)
    x["f_nd"] = (x["dep_status"] == "No actual").astype(int)
    x["f_na"] = ((x["arr_status"] == "No actual") & (x["dep_status"] != "No actual")).astype(int)

    wb = xlsxwriter.Workbook(str(path), {"nan_inf_to_errors": True, "constant_memory": False})
    f = Fmt(wb)
    as_txt = pd.Timestamp(as_of).strftime("%m/%d/%Y %I:%M %p")
    dash = wb.add_worksheet("Dashboard")
    pos = {}

    def h_row(ws, r, c0, crit, g):
        posn = {h: col_letter(c0 + k) + str(r + 1) for k, (h, *_) in enumerate(H_METRICS)}
        for k, (h, tmpl, fn, fk) in enumerate(H_METRICS):
            fo = tmpl.replace("{c}", crit)
            for nm, ad in posn.items():
                fo = fo.replace("{" + nm + "}", ad)
            ws.write_formula(r, c0 + k, fo, getattr(f, fk), _v(fn(g)))

    def h_table(ws, r, title, keys, label, field, fmt):
        ws.write(r, 0, title, f.h2)
        ws.write(r + 1, 0, label, f.hv)
        for k, (h, *_) in enumerate(H_METRICS):
            ws.write(r + 1, 1 + k, h, f.hf)
        ws.set_row(r + 1, 30)
        for i, key in enumerate(keys):
            rr = r + 2 + i
            if isinstance(key, pd.Timestamp):
                ws.write_datetime(rr, 0, key.to_pydatetime(), fmt)
            else:
                ws.write(rr, 0, key, fmt)
            h_row(ws, rr, 1, f"Data[{field}],$A{rr + 1}", x[x[{"Month": "month", "Week": "week"}[field]] == key])
        return r + 2, r + 1 + len(keys)

    ws = wb.add_worksheet("Monthly")
    head(ws, f, "Monthly and weekly trend", 12, as_txt,
         "COUNTIFS / SUMIFS over the Data tab (all trucks scheduled in the period).")
    mkeys = sorted(x["month"].unique())
    pos["m"] = h_table(ws, 3, "By month", mkeys, "Month", "Month", f.txt)
    wkeys = [pd.Timestamp(w) for w in sorted(x["week"].unique())][-26:]
    r = pos["m"][1] + 3
    pos["w"] = h_table(ws, r, "By week (last 26 weeks, Monday start)", wkeys, "Week", "Week", f.date)
    ws.set_column(0, 0, 11)
    ws.set_column(1, 12, 11)
    ws.freeze_panes(3, 1)

    def grid(name, title, field, key, top, metric):
        ws = wb.add_worksheet(name)
        last13 = mkeys[-13:]
        vol = x[x["month"].isin(last13)].groupby(key).size().sort_values(ascending=False).head(top)
        head(ws, f, title, len(last13) + 2, as_txt, f"Top {top} by trucks in the last 13 months. Each cell is a "
             "SUMIFS ratio on Data; blank = no trucks that month.")
        ws.write(3, 0, field, f.hv)
        ws.write(3, 1, "Trucks (13 mo)", f.hf)
        for k, m in enumerate(last13):
            ws.write(3, 2 + k, m, f.hf)
        agg = (x[x["month"].isin(last13)].groupby([key, "month"])
               .agg(n=("f_ld", "size"), ld=("f_ld", "sum"), la=("f_la", "sum"), am=("f_am", "sum")))
        for i, (v, n) in enumerate(vol.items()):
            rr = 4 + i
            ws.write(rr, 0, v)
            ws.write_formula(rr, 1, f'=SUMPRODUCT(COUNTIFS(Data[{field}],$A{rr + 1},Data[Month],$C$4:${col_letter(1 + len(last13))}$4))',
                             f.int_, int(n))
            for k, m in enumerate(last13):
                cl = col_letter(2 + k)
                g = agg.loc[(v, m)] if (v, m) in agg.index else None
                if metric == "otp":
                    fo = (f'=IFERROR(1-SUMIFS(Data[Late Arr],Data[{field}],$A{rr + 1},Data[Month],{cl}$4)/'
                          f'SUMIFS(Data[Arr Measured],Data[{field}],$A{rr + 1},Data[Month],{cl}$4),"")')
                    val = 1 - g["la"] / g["am"] if g is not None and g["am"] else ""
                else:
                    fo = (f'=IFERROR(SUMIFS(Data[Late Dep],Data[{field}],$A{rr + 1},Data[Month],{cl}$4)/'
                          f'COUNTIFS(Data[{field}],$A{rr + 1},Data[Month],{cl}$4),"")')
                    val = g["ld"] / g["n"] if g is not None and g["n"] else ""
                ws.write_formula(rr, 2 + k, fo, f.pct, _v(val))
        if len(vol):
            ws.conditional_format(4, 2, 3 + len(vol), 1 + len(last13), {
                "type": "3_color_scale",
                **({"min_color": RED, "mid_color": "#FFE699", "max_color": "#C6E0B4"} if metric == "otp" else
                   {"min_color": "#FFFFFF", "mid_color": "#FFE699", "max_color": RED})})
        ws.set_column(0, 0, 32)
        ws.set_column(1, 1, 10)
        ws.set_column(2, 2 + len(last13), 8.5)
        ws.freeze_panes(4, 2)

    grid("Vendor OTP", "On-time arrival % by vendor and month", "Vendor", "vendor", 30, "otp")
    grid("Vendor late dep", "Late departure % by vendor and month", "Vendor", "vendor", 30, "dep")
    grid("Origin late dep", "Late departure % by origin and month", "Origin", "origin", 30, "dep")
    grid("Dest OTP", "On-time arrival % by destination and month", "Dest", "dest", 30, "otp")

    ws = wb.add_worksheet("Data")
    ws.write(0, 0, "Data - one row per TLM, last %d months. Statuses use the same rules as the daily report's Data "
                   "formulas (stored as values here to keep the file fast); flags are 1/0." % months, f.sub)
    fm = {"txt": f.txt, "date": f.date, "dtm": f.dtm, "num": f.num}
    rows = [[_cell(v) for v in rec] for rec in x[[c for _, c, _, _ in HIST_COLS]].itertuples(index=False, name=None)]
    ws.add_table(2, 0, 2 + max(len(rows), 1), len(HIST_COLS) - 1, {
        "name": "Data", "data": rows or None, "style": "Table Style Light 1",
        "columns": [{"header": h, "format": fm[fk]} for h, _, _, fk in HIST_COLS]})
    for j, (_, _, w, _) in enumerate(HIST_COLS):
        ws.set_column(j, j, w)
    ws.freeze_panes(3, 2)

    # dashboard
    ws = dash
    ws.hide_gridlines(2)
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    ws.set_column(0, 0, 2)
    ws.set_column(1, 16, 11.5)
    ws.merge_range(0, 1, 0, 12, "Late Trucks - History Dashboard", f.title)
    ws.merge_range(0, 13, 0, 16, f"As of {as_txt}", f.title_r)
    ws.set_row(0, 30)
    ws.write(1, 1, f"{len(x):,} trucks scheduled {start:%m/%d/%Y} - {end:%m/%d/%Y}. Tiles compare the last full month "
                   "with the month before; charts read the Monthly tab.", f.sub)
    mi = {h: 1 + k for k, (h, *_) in enumerate(H_METRICS)}
    a, b = pos["m"]
    full = b - 1 if len(mkeys) > 1 else b       # last full month = second to last row
    prev = full - 1
    tiles = [("Trucks", "Trucks", f.tile_v), ("Late departure %", "Late dep %", f.tile_p),
             ("On-time arrival %", "On-time arr %", f.tile_p), ("Late %", "Late %", f.tile_p),
             ("No ATA", "No ATA", f.tile_r), ("Check date", "Check date", f.tile_v)]
    mtab = x.groupby("month")
    for k, (lab, met, fmt) in enumerate(tiles):
        c = 1 + k * 2
        cl = col_letter(mi[met])
        ws.merge_range(3, c, 3, c + 1, lab, f.tile_h)
        ws.merge_range(4, c, 5, c + 1, "", fmt)
        mk = mkeys[full - a] if 0 <= full - a < len(mkeys) else None
        pk = mkeys[prev - a] if 0 <= prev - a < len(mkeys) else None
        fn = [m for m in H_METRICS if m[0] == met][0][2]
        ws.write_formula(4, c, f"=Monthly!{cl}{full + 1}", fmt, _v(fn(mtab.get_group(mk)) if mk else ""))
        is_pct = "%" in met
        ws.merge_range(6, c, 6, c + 1, "", f.sub)
        fmt_t = "+0.0%;-0.0%" if is_pct else "+#,##0;-#,##0"
        cur_v = fn(mtab.get_group(mk)) if mk else ""
        prev_v = fn(mtab.get_group(pk)) if pk else ""
        dv = ""
        if cur_v != "" and prev_v != "":
            d = cur_v - prev_v
            dv = "vs prior month: " + (f"{d:+.1%}" if is_pct else f"{d:+,.0f}")
        ws.write_formula(6, c, f'="vs prior month: "&IFERROR(TEXT(Monthly!{cl}{full + 1}-Monthly!{cl}{prev + 1},'
                               f'"{fmt_t}"),"")', f.sub, dv)
    if mkeys:
        ws.write(2, 1, f"Last full month: {mkeys[full - a] if 0 <= full - a < len(mkeys) else ''}", f.sub)

    def mchart(title, rows, series, combo=None, pct=True):
        a, b = rows
        ch = wb.add_chart({"type": "column" if combo else "line"})
        cat = ["Monthly", a - 1 + 0, 0, b - 1, 0]
        if combo:
            ch.add_series({"name": combo[0], "categories": cat, "values": ["Monthly", a - 1, combo[1], b - 1, combo[1]],
                           "fill": {"color": "#BDD7EE"}, "gap": 50})
            ln = wb.add_chart({"type": "line"})
            for nm, cc, color in series:
                ln.add_series({"name": nm, "categories": cat, "values": ["Monthly", a - 1, cc, b - 1, cc],
                               "y2_axis": True, "line": {"color": color, "width": 2.25}})
            ln.set_y2_axis({"num_format": "0%", "min": 0, "max": 1})
            ch.combine(ln)
        else:
            for nm, cc, color in series:
                ch.add_series({"name": nm, "categories": cat, "values": ["Monthly", a - 1, cc, b - 1, cc],
                               "line": {"color": color, "width": 2.25}})
            ch.set_y_axis({"num_format": "0%" if pct else "0", "min": 0})
        ch.set_title({"name": title, "name_font": {"size": 11, "bold": True}})
        ch.set_legend({"position": "bottom"})
        ch.set_x_axis({"num_font": {"rotation": -45, "size": 8}})
        ch.set_chartarea({"border": {"none": True}})
        return ch

    m_rows = (pos["m"][0] + 1, pos["m"][1] + 1)
    ws.insert_chart(8, 1, mchart("Monthly trucks, late departure % and on-time arrival %", m_rows,
                                 [("Late dep %", mi["Late dep %"], RED), ("On-time arr %", mi["On-time arr %"], GREEN)],
                                 combo=("Trucks", mi["Trucks"])), {"x_scale": 1.55, "y_scale": 1.2})
    ws.insert_chart(8, 9, mchart("Monthly follow-ups: No ATD, No ATA, Check date", m_rows,
                                 [("No ATD", mi["No ATD"], "#7F7F7F"), ("No ATA", mi["No ATA"], "#F28E2B"),
                                  ("Check date", mi["Check date"], AMBER)], pct=False), {"x_scale": 1.25, "y_scale": 1.2})
    w_rows = (pos["w"][0] + 1, pos["w"][1] + 1)
    ch = wb.add_chart({"type": "line"})
    for nm, cc, color in [("Late dep %", mi["Late dep %"], RED), ("On-time arr %", mi["On-time arr %"], GREEN)]:
        ch.add_series({"name": nm, "categories": ["Monthly", w_rows[0] - 1, 0, w_rows[1] - 1, 0],
                       "values": ["Monthly", w_rows[0] - 1, cc, w_rows[1] - 1, cc],
                       "line": {"color": color, "width": 2.25}})
    ch.set_title({"name": "Weekly late departure % and on-time arrival % (26 weeks)",
                  "name_font": {"size": 11, "bold": True}})
    ch.set_y_axis({"num_format": "0%", "min": 0})
    ch.set_x_axis({"num_format": "m/d", "num_font": {"rotation": -45, "size": 8}})
    ch.set_legend({"position": "bottom"})
    ch.set_chartarea({"border": {"none": True}})
    ws.insert_chart(27, 1, ch, {"x_scale": 1.55, "y_scale": 1.1})
    ws.activate()
    wb.close()
    return path
